"""2.9.107: the CNSS contribution forms as editable Excel, in the layout of the printed forms (RTL, A4, one page):

- CNSS 190A - contributions due for a month or a quarter (جدول الإشتراكات المستحقة)
- CNSS 351  - contributions due from the annual nominative declaration (جدول الإشتراكات المتوجبة والناتجة عن التصريح الإسمي السنوي)
- CNSS 386  - annual nominative declaration (التصريح الإسمي السنوي) + the attached employee / month schedule

Filled from Settings > Company and the posted payroll (the same figures as the NSSF report). Amounts are real numbers and the
totals are Excel formulas: contribution = wages x rate, totals, net due after the family allowances paid, the annual
settlement. Light-yellow cells stay editable and every formula follows a changed value."""
from __future__ import annotations

import re
from datetime import date

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from payroll_excel_forms import INPUT, Form, center, right, thin

MONEY = '#,##0;-#,##0;""'
TOTAL = "#,##0;-#,##0;0"  # computed cells show 0 rather than staying empty
RATE = "0.00%"
GREY = "7F7F7F"
MONTHS_AR = ("كانون الثاني", "شـبـاط", "آذار", "نيـســان", "أيـــار", "حزيران", "تمـــوز", "آب", "أيـلـول", "تشـرين الأول",
             "تشـرين الثاني", "كانون الأول")
BRANCHES = (("sick", "المرض والأمومة"), ("eos", "تعويض نهاية الخدمة"), ("family", "التعويضات العائلية"))
DEFAULT_RATES = {"sick": 0.11, "eos": 0.085, "family": 0.06}  # used only when there is no payroll and no saved rule


# ---------------------------------------------------------------------------------------------------- the figures
def _number(value):
    if value in (None, ""): return None
    try: return float(str(value).replace(",", ""))
    except (TypeError, ValueError): return None


def _percent(text):
    match = re.search(r"-?\d+(?:\.\d+)?", str(text or ""))
    return float(match.group()) / 100 if match else None


def detail_rows(report):
    """Employee / month lines of the NSSF report: [nssf no, name, MM-YYYY, salary, sick base, employee, employer sick,
    family base, family, eos base, eos, total, allowance, net]."""
    detail = (report or {}).get("monthly_detail")
    if isinstance(detail, dict): rows = detail.get("rows") or []
    else:
        sections = (report or {}).get("sections") or []
        rows = sections[3].get("rows") if len(sections) > 3 and isinstance(sections[3], dict) else []
    return [row for row in rows if isinstance(row, (list, tuple)) and len(row) >= 11 and re.fullmatch(r"\d{2}-\d{4}", str(row[2] or ""))]


def saved_rates(report):
    """The rates in force (last rule of the period) from the report: sick = employee + employer share."""
    sections = (report or {}).get("sections") or []
    rules = sections[2].get("rows") if len(sections) > 2 and isinstance(sections[2], dict) else []
    rules = [row for row in rules or [] if isinstance(row, (list, tuple)) and len(row) >= 7 and _percent(row[3]) is not None]
    if not rules: return dict(DEFAULT_RATES)
    last = rules[-1]
    return {"sick": (_percent(last[3]) or 0) + (_percent(last[4]) or 0), "family": _percent(last[5]) or 0, "eos": _percent(last[6]) or 0}


def branch_figures(report):
    """Per branch: number of employees, wages (capped bases), rate and contribution, from the NSSF report.
    The rate is the contribution divided by the wages, so wages x rate gives back the payroll amount exactly."""
    summary = (report or {}).get("summary") or {}
    rates = saved_rates(report)
    rows = detail_rows(report)
    result = {}
    for branch, base_key, due_keys, column in (("sick", "sick_base", ("employee", "employer_sick"), 4),
                                               ("eos", "eos_base", ("eos",), 9), ("family", "family_base", ("family",), 7)):
        wages = _number(summary.get(base_key)) or 0.0
        due = sum(_number(summary.get(key)) or 0.0 for key in due_keys)
        count = len({(row[0], row[1]) for row in rows if (_number(row[column]) or 0) > 0})
        result[branch] = {"count": count, "wages": wages, "due": due, "rate": (due / wages) if wages else rates.get(branch, DEFAULT_RATES[branch])}
    result["allowance"] = _number(summary.get("allowance")) or 0.0
    return result


def employee_months(report):
    """{(nssf no, name): {month: {"sick", "family", "eos"}}} from the yearly NSSF report, for the 386 schedule."""
    employees = {}
    for row in detail_rows(report):
        month = int(str(row[2])[:2])
        values = employees.setdefault((str(row[0] or "").strip(), str(row[1] or "").strip()), {}).setdefault(month, {"sick": 0.0, "family": 0.0, "eos": 0.0})
        for branch, column in (("sick", 4), ("family", 7), ("eos", 9)): values[branch] += _number(row[column]) or 0.0
    return dict(sorted(employees.items(), key=lambda item: item[0][1]))


def monthly_bases(report):
    totals = {month: {"sick": 0.0, "family": 0.0, "eos": 0.0} for month in range(1, 13)}
    for months in employee_months(report).values():
        for month, values in months.items():
            for branch in totals[month]: totals[month][branch] += values[branch]
    return totals


def filed_bases(filed_wages, report):
    """Section (a) of the 386: the wages declared each month. Saved filed wages when present; otherwise the posted payroll
    of that month (the wages Saber declared). Returns ({month: {branch: amount}}, [months taken from payroll])."""
    actual = monthly_bases(report)
    saved = {int(row.get("month")): row for row in filed_wages or [] if row.get("month") not in (None, "")}
    result, from_payroll = {}, []
    for month in range(1, 13):
        row = saved.get(month) or {}
        values = {branch: _number(row.get(key)) for branch, key in (("sick", "sickness_wages"), ("family", "family_wages"), ("eos", "end_service_wages"))}
        if any(value is None for value in values.values()):
            values = {branch: (values[branch] if values[branch] is not None else actual[month][branch]) for branch in values}
            if any(actual[month].values()): from_payroll.append(month)
        result[month] = values
    return result, from_payroll


# ---------------------------------------------------------------------------------------------------- shared drawing
def _money(f, r, c1, c2, value=None, formula=None, bold=False, editable=True, r2=None):
    cell = f.field(r, c1, c2, r2=r2, box=True) if editable else f.merge(r, c1, c2, r2)
    if formula: cell.value = formula
    elif value not in (None, ""): cell.value = round(float(value))
    cell.number_format = MONEY if editable else TOTAL; cell.alignment = Alignment(horizontal="center", vertical="center", shrink_to_fit=True)
    cell.font = Font(name="Arial", size=10, bold=bold, color="1F3B73")
    if not editable: _box(f, r, (r2 or r), c1, c2)
    return cell


def _box(f, r1, r2, c1, c2, style="thin"):
    s = Side(style=style, color="000000")
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            x = f.ws.cell(r, c); b = x.border
            x.border = Border(right=s if c == c2 else b.right, left=s if c == c1 else b.left,
                              top=s if r == r1 else b.top, bottom=s if r == r2 else b.bottom)


def _head(f, r, c1, c2, text, size=9, r2=None, rotate=False, fill="F2F2F2"):
    align = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2, text_rotation=90 if rotate else 0)
    cell = f.label(r, c1, c2, text, size, True, align=align, r2=r2)
    _box(f, r, r2 or r, c1, c2)
    if fill: f.fill(r, r2 or r, c1, c2, fill)
    return cell


def _note(f, r, c1, c2, text, size=7, r2=None):
    cell = f.label(r, c1, c2, text, size, align=Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2), r2=r2)
    cell.font = Font(name="Arial", size=size, color=GREY); _box(f, r, r2 or r, c1, c2)
    return cell


def _landscape(f, path, last_row):
    f.page(last_row)
    f.ws.page_setup.orientation = "landscape"
    if path: f.wb.save(path)
    return f.wb


def _payment_methods(f, r, c1, c2, method):
    """Payment method as printed: 3 cheque, 2 bank transfer, 1 cash, 0 unpaid; the chosen one gets an X."""
    method = str(method or "").lower()
    chosen = "3" if "cheq" in method or "شك" in method else "2" if "transfer" in method or "تحويل" in method else "1" if "cash" in method or "نقد" in method \
        else "0" if "unpaid" in method or "غير" in method else ""
    for i, (text, code) in enumerate((("شـك", "3"), ("تحويل مصرفي", "2"), ("نقداً", "1"), ("غير مدفوع", "0"))):
        f.label(r + i, c1, c2 - 1, text, 8)
        box = f.ws.cell(r + i, c2); box.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        box.fill = PatternFill("solid", fgColor=INPUT); box.alignment = center; box.font = Font(name="Arial", size=9, bold=True, color="1F3B73")
        box.value = "X" if chosen == code else code


# ---------------------------------------------------------------------------------------------------- 190A and 351
def _contributions(f, data, yearly):
    """The body shared by CNSS 190A and CNSS 351 (contributions table, totals, signatures)."""
    N = f.N
    f.label(1, 1, 9, "الصندوق الوطني", 14, True); f.label(2, 1, 9, "للضمان الإجتماعي", 14, True); f.label(3, 1, 9, "بيروت - لبنان", 11, True)
    if yearly:
        f.label(1, 11, 26, "جدول الإشتراكات المتوجبة والناتجة عن التصريح الإسمي السنوي لعام", 10, True)
        f.field(1, 27, 31, key="year"); f.ws.cell(1, 27).alignment = center
        f.label(2, 11, 22, "والمتوجب دفعها حتى تاريخ", 10, True); f.field(2, 23, 31, key="pay_until")
    else:
        f.label(1, 11, 22, "جدول الإشتراكات المستحقة عن", 10, True); f.field(1, 23, 33, key="period_text")
        f.label(2, 11, N, "المتوجب تسديدها خلال مهلة خمسة عشر يوماً إعتباراً من نهاية الشهر الإستحقاق أو من تاريخ الإستلام", 9)
    f.heights((1, 2, 3), 20)

    # employer box (right) | fund box (left)
    r = 4; _head(f, r, 1, 16, "حقل مخصص لأصحاب العمل", 10, fill=None); _head(f, r, 17, N, "حقل مخصص للصندوق", 10, fill=None)
    f.label(5, 1, 4, "إسم المؤسسـة :", 9, True); f.field(5, 5, 16, r2=6, key="company_name")
    f.label(8, 1, 4, "رقم المؤسسـة :", 9, True); f.boxes(8, 6, 10, key="company_nssf")
    _box(f, 5, 9, 1, 16)
    f.label(5, 17, 22, "رقم المستند", 8, True, align=center); f.field(6, 17, 22, key="receipt_number")
    f.label(7, 17, 22, "تاريخ الدفع (يوم / شهر / سنة)", 8, True, align=center); f.field(8, 17, 22, key="payment_date")
    _box(f, 5, 9, 17, 22)
    f.label(5, 23, 27, "طريقة دفع الإشتراكات", 8, True, align=center); _payment_methods(f, 6, 23, 27, data.get("payment_method")); _box(f, 5, 9, 23, 27)
    if yearly:
        f.label(5, 28, 31, "المدة", 8, True, align=center); f.label(6, 28, 28, "من", 8); f.field(6, 29, 31, key="date_from"); f.label(7, 28, 28, "إلى", 8); f.field(7, 29, 31, key="date_to")
        _box(f, 5, 9, 28, 31)
        f.label(5, 32, N, "مؤسســـة", 8, True, align=center)
        for i, (text, value) in enumerate((("فصليــة", "quarterly"), ("شهريـة", "monthly"))):
            f.label(6 + i, 32, 35, text, 8, True); box = f.ws.cell(6 + i, N); box.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            box.fill = PatternFill("solid", fgColor=INPUT); box.alignment = center; box.value = "X" if data.get("declaration") == value else None
        _box(f, 5, 9, 32, N)
    else:
        f.label(5, 28, 32, "المدة", 8, True, align=center); f.label(6, 28, 29, "من", 8); f.field(6, 30, 32, key="date_from"); f.label(7, 28, 29, "إلى", 8); f.field(7, 30, 32, key="date_to")
        _box(f, 5, 9, 28, 32)
        f.label(5, 33, N, "ملاحظات", 8, True, align=center); f.field(6, 33, N, r2=9); _box(f, 5, 9, 33, N)
    f.heights(range(4, 10), 15)

    # the contributions table
    r = 11
    _head(f, r, 1, 2, "فرع", r2=r + 1); _head(f, r, 3, 4, "عدد الأجراء", 8, r2=r + 1); _head(f, r, 5, 9, "الأجور ولواحقها", r2=r + 1)
    _head(f, r, 10, 11, "المعدل", 8, r2=r + 1); _head(f, r, 12, 16, "الإشتراكات المستحقة", r2=r + 1)
    _head(f, r, 17, 23, "الإشتراكات المدفوعة" if yearly else "الإشتراكات"); _head(f, r + 1, 17, 20, "المدفوع", 8); _head(f, r + 1, 21, 23, "الرمز", 8)
    _head(f, r, 24, 34, "زيادات التأخير"); _head(f, r + 1, 24, 25, "عدد أيام التأخير", 7); _head(f, r + 1, 26, 29, "المتوجب", 8); _head(f, r + 1, 30, 34, "المدفوع", 8)
    _head(f, r, 35, N, "مرجع/فئة/نوع/فرع" if not yearly else "الفرع", 7, r2=r + 1)
    f.heights((r, r + 1), 22)
    refs = {"sick": "3/0/1/0", "eos": "1/0/1/0", "family": "2/0/1/0"} if not yearly else {"sick": "3", "eos": "1", "family": "2"}
    first = r + 2
    for i, (branch, name) in enumerate(BRANCHES):
        row = first + i; values = data.get(branch) or {}
        _head(f, row, 1, 2, name, 7, rotate=False, fill="F2F2F2")
        count = f.field(row, 3, 4, box=True); count.value = values.get("count") or None; count.alignment = center; count.number_format = "0"
        _money(f, row, 5, 9, values.get("wages"))
        rate = f.field(row, 10, 11, box=True); rate.value = values.get("rate"); rate.number_format = RATE; rate.alignment = center
        _money(f, row, 12, 16, formula=f'=IF(E{row}="","",ROUND(E{row}*J{row},0))', bold=True)
        _money(f, row, 17, 20, values.get("paid"))
        _note(f, row, 21, 23, "01 / 11")
        days = f.field(row, 24, 25, box=True); days.alignment = center; days.number_format = "0"
        _money(f, row, 26, 29, formula=f'=IF(OR(X{row}="",L{row}=""),"",ROUND(L{row}*X{row}/2000,0))')
        _money(f, row, 30, 34)
        _note(f, row, 35, N, refs[branch], 8)
    f.heights(range(first, first + 3), 34)
    last_branch = first + 2

    # totals
    r = last_branch + 1
    f.label(r, 1, 11, "مجموع الإشتراكات المستحقة", 9, True); _box(f, r, r, 1, 11)
    _money(f, r, 12, 16, formula=f"=SUM(L{first}:L{last_branch})", bold=True, editable=False)
    _money(f, r, 17, 20, formula=f"=SUM(Q{first}:Q{last_branch})", editable=False)
    _money(f, r, 26, 29, formula=f"=SUM(Z{first}:Z{last_branch})", editable=False)
    _money(f, r, 30, 34, formula=f"=SUM(AD{first}:AD{last_branch})", editable=False)
    total_row = r
    r += 1; f.label(r, 1, 11, "التعويضات العائلية المدفوعة", 9, True); _box(f, r, r, 1, 11); _money(f, r, 12, 16, data.get("allowance"))
    allowance_row = r
    r += 1; f.label(r, 1, 11, "الزيادة المتوجبة على الصندوق / الباقي المتوجب للصندوق", 9, True); _box(f, r, r, 1, 11)
    _money(f, r, 12, 16, formula=f"=L{total_row}-L{allowance_row}", bold=True, editable=False)
    net_row = r
    f.label(total_row + 1, 17, 25, "مجموع المبالغ المدفوعة", 9, True, align=center); _box(f, total_row + 1, total_row + 1, 17, 25)
    _money(f, total_row + 1, 26, 34, formula=f"=SUM(Q{total_row},AD{total_row})", bold=True, editable=False)
    f.label(total_row + 2, 17, 25, "الباقي المتوجب دفعه (مع زيادات التأخير)", 9, True, align=center); _box(f, total_row + 2, total_row + 2, 17, 25)
    _money(f, total_row + 2, 26, 34, formula=f"=L{net_row}+Z{total_row}-Q{total_row}-AD{total_row}", bold=True, editable=False)
    f.heights(range(total_row, net_row + 1), 20)

    # declaration and signatures
    r = net_row + 2
    f.label(r, 1, 22, "إن رب العمل الموقع أدناه يشهد على مسؤوليته أن الأجور وكافة المعلومات المصرح بها أعلاه هي مطابقة للحقيقة والواقع", 8, True, r2=r + 1)
    f.label(r, 24, 27, "المصفي", 9, True); f.field(r, 28, N)
    f.label(r + 1, 24, 27, "المراقب", 9, True); f.field(r + 1, 28, N)
    f.label(r + 2, 24, 27, "المحاسب", 9, True); f.field(r + 2, 28, N)
    f.label(r + 2, 1, 2, "في", 9, True); f.field(r + 2, 3, 8, key="today"); f.label(r + 2, 10, 14, "ختم المؤسسة", 9, True, align=center)
    f.label(r + 2, 16, 22, "توقيع رب العمل", 9, True, align=center)
    f.field(r + 3, 10, 14, r2=r + 5, box=True); f.field(r + 3, 16, 22, r2=r + 5, box=True)
    f.heights(range(r, r + 6), 18)
    r += 6
    text = ("ملاحظة هامة : يتوجب إعادة هذا الجدول مع جدول التصريح الإسمي السنوي المرفق" if yearly else
            "ملاحظة هامة : يتوجب إعادة هذا الجدول مع جدول التعويضات العائلية المرفق تحت طائلة عدم الأخذ بقيمة التعويضات العائلية المدفوعة من قبلكم")
    f.label(r, 1, N, text, 8, True); _box(f, r, r, 1, N)
    f.label(r + 1, 1, 8, "CNSS 351" if yearly else "CNSS 190A", 9, True)
    f.frame(4, r)
    return r + 1


def build_cnss_190a(data=None, path=None):
    f = Form("CNSS 190A", cols=36, width=3.7, data=data)
    last = _contributions(f, data or {}, yearly=False)
    return _landscape(f, path, last)


def build_cnss_351(data=None, path=None):
    f = Form("CNSS 351", cols=36, width=3.7, data=data)
    last = _contributions(f, data or {}, yearly=True)
    return _landscape(f, path, last)


# ---------------------------------------------------------------------------------------------------- 386
def build_cnss_386(data=None, path=None):
    """Annual nominative declaration: section (a) wages declared each month, section (b) the annual settlement
    (wages of the attached schedule - wages declared) x rates, and the attached employee / month schedule sheet."""
    data = data or {}
    f = Form("CNSS 386", cols=24, width=4.6, data=data); N = f.N
    schedule_sheet = "الجدول الملحق"
    f.label(1, 1, 10, "الصــــــندوق الوطني", 15, True); f.label(2, 1, 10, "للضمان الإجتماعي", 15, True)
    f.label(4, 1, 10, "التصريح الإسمي السنوي", 14, True); f.label(5, 1, 3, "لســـنة", 11, True); f.field(5, 4, 7, key="year")
    f.label(1, 11, 14, "جانب صاحب العمل", 10, True, align=center, r2=3)
    f.label(1, 15, 16, "الرقم", 10, True); f.boxes(1, 17, 8, key="company_nssf")
    f.field(2, 15, N, r2=3, key="company_name"); f.field(4, 15, N, key="company_address"); _box(f, 1, 4, 15, N)
    f.heights(range(1, 6), 20)
    f.label(6, 1, N, "بناءً على المادة 80 (فقرة 5) من قانون الضمان الإجتماعي المعدل بموجب المرسوم الإشتراعي رقم 73/116 وبناءً على نظام الإشتراكات لدى الصندوق. "
            "ندعوكم إلى كشف المعلومات المطلوبة أدناه وإعادة هذا التصريح والجدول الملحق به مع المعلومات المطلوبة ودفع الرصيد المتوجب خلال مهلة أقصاها 31 آذار",
            9, r2=8)

    # (a) wages declared during the year
    r = 10; f.label(r, 1, N, "(أ) – مراجعة الأجور التي دفعت عنها الإشتراكات خلال السـنة:", 10, True, True)
    r = 11; _head(f, r, 1, 4, "المدة"); _head(f, r, 5, 8, "فرع نهاية الخدمة"); _head(f, r, 9, 12, "فرع التعويضات العائلية"); _head(f, r, 13, 16, "فرع المرض والأمومة")
    columns = {"eos": (5, 8, "E"), "family": (9, 12, "I"), "sick": (13, 16, "M")}
    filed = data.get("filed") or {}
    first = r + 1
    for month in range(1, 13):
        row = first + month - 1; f.label(row, 1, 4, MONTHS_AR[month - 1], 9); _box(f, row, row, 1, 4)
        for branch, (c1, c2, _letter) in columns.items(): _money(f, row, c1, c2, (filed.get(month) or {}).get(branch))
    diff_row = first + 12; f.label(diff_row, 1, 4, "فـــروقات", 9, True); _box(f, diff_row, diff_row, 1, 4)
    for branch, (c1, c2, _letter) in columns.items(): _money(f, diff_row, c1, c2)
    sum_row = diff_row + 1; f.label(sum_row, 1, 4, "المجاميع", 9, True); _box(f, sum_row, sum_row, 1, 4)
    for branch, (c1, c2, letter) in columns.items(): _money(f, sum_row, c1, c2, formula=f"=SUM({letter}{first}:{letter}{diff_row})", bold=True, editable=False)
    for caption, (c1, c2) in (("(حقل 1)", (5, 8)), ("(حقل 2)", (9, 12)), ("(حقل 3)", (13, 16))): f.label(sum_row + 1, c1, c2, caption, 8, align=center)
    f.heights(range(r, sum_row + 2), 16)
    for i, title in enumerate(("المصفي", "المراقب", "المحاسب")):
        f.label(first + 3 + i * 3, 18, 19, title, 10, True); f.field(first + 3 + i * 3, 20, N)

    # (b) annual settlement
    r = sum_row + 3
    f.label(r, 1, 8, "(ب) – إشتراكات التسوية السنوية المتوجب تسديدها:", 10, True)
    _head(f, r, 9, 13, "فرع نهاية الخدمة"); _head(f, r, 14, 18, "فرع التعويضات العائلية"); _head(f, r, 19, 23, "فرع المرض والأمومة")
    b_cols = {"eos": (9, 13, "I", "E"), "family": (14, 18, "N", "I"), "sick": (19, 23, "S", "M")}
    r1 = r + 1; f.label(r1, 1, 8, "مجموع الأجور وفقاً للجدول الملحق بهذا التصريح", 9)
    for branch, (c1, c2, _l, _a) in b_cols.items():
        _money(f, r1, c1, c2, formula=f"=__SCHEDULE_{branch}", editable=False)  # set by _schedule once its rows are known
    f.label(r1 + 1, 1, 8, "(نــاقص)", 9, True, align=center)
    for caption, (c1, c2) in (("(حقل 4)", (9, 13)), ("(حقل 5)", (14, 18)), ("(حقل 6)", (19, 23))): f.label(r1 + 1, c1, c2, caption, 8, align=center)
    r2 = r1 + 2; f.label(r2, 1, 8, "مجموع الأجور كما هو مبين في الحقل ( أ ) أعلاه", 9)
    for branch, (c1, c2, _l, a_letter) in b_cols.items(): _money(f, r2, c1, c2, formula=f"={a_letter}{sum_row}", editable=False)
    f.label(r2 + 1, 1, 8, "(يســاوي)", 9, True, align=center)
    for caption, (c1, c2) in (("(حقل 1)", (9, 13)), ("(حقل 2)", (14, 18)), ("(حقل 3)", (19, 23))): f.label(r2 + 1, c1, c2, caption, 8, align=center)
    r3 = r2 + 2; f.label(r3, 1, 8, "أجور التســوية الســنوية:", 9)
    for branch, (c1, c2, letter, _a) in b_cols.items(): _money(f, r3, c1, c2, formula=f"={letter}{r1}-{letter}{r2}", editable=False)
    r4 = r3 + 1; f.label(r4, 1, 8, "المعدل", 9, align=center)
    rates = data.get("rates") or DEFAULT_RATES
    for branch, (c1, c2, _l, _a) in b_cols.items():
        cell = f.field(r4, c1, c2, box=True); cell.value = rates.get(branch); cell.number_format = RATE; cell.alignment = center
    r5 = r4 + 1; f.label(r5, 1, 8, "إشتراكات التسوية", 9, True)
    for branch, (c1, c2, letter, _a) in b_cols.items(): _money(f, r5, c1, c2, formula=f"=ROUND({letter}{r3}*{letter}{r4},0)", bold=True, editable=False)
    f.label(r5, 24, 24, "=", 11, True, align=center)
    f.heights(range(r, r5 + 1), 18)
    r6 = r5 + 2; f.label(r6, 1, 12, "مجموع إشتراكات التسوية المتوجب دفعها", 10, True)
    _money(f, r6, 13, 18, formula=f"=SUM(I{r5}:W{r5})", bold=True, editable=False); f.heights((r6,), 22)

    # (c) declaration
    r = r6 + 2; f.label(r, 1, N, "(ج) – (مطبوعة ملحقة بالتصريح الإسمي السنوي)", 10, True)
    f.label(r + 1, 1, 4, "أنا الموقع أدناه", 9); f.field(r + 1, 5, 14, key="signatory"); f.label(r + 1, 15, 16, "بصفتي", 9); f.field(r + 1, 17, N, key="signatory_title")
    f.label(r + 2, 1, N, "أشهد أن جميع المعلومات التي أعطيتها على هذا التصريح والجدول الملحق به، مطابقة للحقيقة بعد إجراء جميع التصحيحات والتعديلات "
            "والإضافات اللازمة على الجدول الملحق المذكور.", 9, r2=r + 3)
    r += 5; f.label(r, 1, 5, "ختم المؤســســة", 10, True); f.label(r, 12, 13, "في", 9); f.field(r, 14, 19, key="today")
    f.label(r + 1, 12, 14, "التوقيع", 9); f.field(r + 1, 15, N); f.field(r + 1, 1, 5, r2=r + 3, box=True)
    f.label(r + 4, 1, 4, "CNSS386", 9, True)
    last = r + 4
    f.page(last)

    _schedule(f, data, schedule_sheet)
    if path: f.wb.save(path)
    return f.wb


def _schedule(f, data, title):
    """The attached schedule: one row per employee and branch, wages of each month, totals by formula."""
    ws = f.wb.create_sheet(title); ws.sheet_view.rightToLeft = True; ws.sheet_view.showGridLines = False
    f.ws = ws
    widths = {"A": 12, "B": 24, "C": 14, "P": 14}
    from openpyxl.utils import get_column_letter as L
    for c in range(1, 17): ws.column_dimensions[L(c)].width = widths.get(L(c), 11)
    bold = Font(name="Arial", size=10, bold=True)
    ws.merge_cells("A1:P1"); ws["A1"].value = f"الجدول الملحق بالتصريح الإسمي السنوي لسنة {data.get('year') or ''}"
    ws["A1"].font = Font(name="Arial", size=13, bold=True); ws["A1"].alignment = center
    ws.merge_cells("A2:H2"); ws["A2"].value = f"إسم المؤسسة: {data.get('company_name') or ''}"; ws["A2"].font = bold; ws["A2"].alignment = right
    ws.merge_cells("I2:P2"); ws["I2"].value = f"رقم المؤسسة في الصندوق: {data.get('company_nssf') or ''}"; ws["I2"].font = bold; ws["I2"].alignment = right
    headers = ["رقم الأجير في الصندوق", "إسم الأجير وشهرته", "الفرع"] + list(MONTHS_AR) + ["المجموع"]
    for c, text in enumerate(headers, 1):
        cell = ws.cell(4, c, text); cell.font = Font(name="Arial", size=9, bold=True); cell.alignment = center
        cell.fill = PatternFill("solid", fgColor="F2F2F2"); cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.row_dimensions[4].height = 30
    row = 5
    branch_names = {"sick": "المرض والأمومة", "family": "التعويضات العائلية", "eos": "نهاية الخدمة"}
    employees = data.get("employees") or {}
    for (number, name), months in employees.items():
        for branch in ("sick", "family", "eos"):
            values = [number, name, branch_names[branch]] + [round((months.get(m) or {}).get(branch) or 0) or None for m in range(1, 13)]
            for c, value in enumerate(values, 1):
                cell = ws.cell(row, c, value); cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                cell.font = Font(name="Arial", size=9, color="1F3B73" if c > 3 else "000000"); cell.alignment = center if c != 2 else right
                if c > 3: cell.number_format = MONEY; cell.fill = PatternFill("solid", fgColor=INPUT)
            total = ws.cell(row, 16, f"=SUM(D{row}:O{row})"); total.number_format = MONEY; total.font = Font(name="Arial", size=9, bold=True)
            total.border = Border(left=thin, right=thin, top=thin, bottom=thin); total.alignment = center
            row += 1
    extra = 6  # empty rows for employees added by hand
    for _ in range(extra):
        for c in range(1, 17):
            cell = ws.cell(row, c); cell.border = Border(left=thin, right=thin, top=thin, bottom=thin); cell.fill = PatternFill("solid", fgColor=INPUT)
            if c > 3: cell.number_format = MONEY
        ws.cell(row, 16, f"=SUM(D{row}:O{row})").number_format = MONEY
        row += 1
    last_data = row - 1
    totals = {}
    for branch in ("sick", "family", "eos"):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        ws.cell(row, 1, f"مجموع فرع {branch_names[branch]}").font = bold; ws.cell(row, 1).alignment = right
        for c in range(4, 17):
            L_ = L(c); cell = ws.cell(row, c, f'=SUMIF($C$5:$C${last_data},"{branch_names[branch]}",{L_}5:{L_}{last_data})')
            cell.number_format = MONEY; cell.font = bold; cell.alignment = center
        for c in range(1, 17): ws.cell(row, c).border = Border(left=thin, right=thin, top=thin, bottom=thin); ws.cell(row, c).fill = PatternFill("solid", fgColor="F2F2F2")
        totals[branch] = row; row += 1
    ws.cell(row + 1, 1, "الفرع: اكتب الإسم كما هو (المرض والأمومة / التعويضات العائلية / نهاية الخدمة) ليدخل في المجموع.").font = Font(name="Arial", size=8, color=GREY)
    ws.print_area = f"A1:P{row}"; ws.print_title_rows = "4:4"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4; ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.freeze_panes = "D5"
    # the 386 refers to these totals: fill the references now that the rows are known
    main = f.wb.worksheets[0]
    for r in range(1, main.max_row + 1):
        for c in range(1, main.max_column + 1):
            cell = main.cell(r, c)
            if isinstance(cell.value, str) and cell.value.startswith("=__SCHEDULE_"):
                branch = cell.value[len("=__SCHEDULE_"):]
                cell.value = f"='{ws.title}'!P{totals[branch]}"
    return totals
