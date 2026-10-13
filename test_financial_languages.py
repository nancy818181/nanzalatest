"""2.9.109: financial statements, notes and auditor's report in English, French and Arabic."""
import tests_setup  # noqa: F401
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

import financial_statements as fs
import report_export as rx
from database import Database


class FinancialLanguagesTest(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory(ignore_cleanup_errors=True); self.addCleanup(self.folder.cleanup)
        self.db = Database(Path(self.folder.name) / "c.db"); self.db.initialize("secret12345"); self.addCleanup(self.db.release)
        self.db.save_settings({"company_name": "Alpha SARL", "company_name_ar": "ألفا ش.م.م."}, 1)
        self.db.import_invoice({"invoice_number": "S-1", "invoice_date": "10-03-2026", "party_name": "Client", "kind": "sale", "currency": "USD",
                                "subtotal": 1000, "vat": 110, "total": 1110}, 1)

    def pack(self, language):
        return fs.build({2026: self.db}, {"years": "2026", "basis": "USD", "language": language})

    def test_same_figures_in_every_language(self):
        figures = {}
        for language in ("en", "fr", "ar"):
            result = self.pack(language)
            figures[language] = [[str(v) for v in row[1:]] for s in result["sections"] if s.get("kind") in ("position", "profit", "cash") for row in s["rows"]]
        self.assertEqual(figures["en"], figures["fr"]); self.assertEqual(figures["en"], figures["ar"])

    def test_texts_are_translated(self):
        fr, ar = self.pack("fr"), self.pack("ar")
        self.assertIn("ÉTAT DE LA SITUATION FINANCIÈRE", next(s["heading"] for s in fr["sections"] if s["kind"] == "position"))
        self.assertIn("بيان المركز المالي", next(s["heading"] for s in ar["sections"] if s["kind"] == "position"))
        self.assertEqual(ar["meta"][0], "ألفا ش.م.م.")                  # the Arabic company name in the Arabic pack
        self.assertTrue(any("Nous avons effectué l'audit" in r[0] for r in fr["sections"][0]["rows"]))
        self.assertTrue(any("لقد دققنا البيانات المالية" in r[0] for r in ar["sections"][0]["rows"]))
        self.assertEqual(self.pack("en")["title"], "Financial Statements, Notes and Audit Report")

    def test_a_text_the_user_rewrote_is_kept(self):
        fs.save_config(self.db, {"notes": {"Events and going concern": "Texte écrit par l'auditeur."}}, 1)
        rows = [r[0] for s in self.pack("ar")["sections"] for r in s["rows"] if s.get("narrative")]
        self.assertIn("Texte écrit par l'auditeur.", rows)

    def test_excel_pack_in_french_and_arabic(self):
        for language, cover in (("fr", "Couverture"), ("ar", "الغلاف")):
            result = self.pack(language); path = Path(self.folder.name) / f"fs_{language}.xlsx"
            rx.export_sections_excel(str(path), result["title"], result["meta"], result["sections"])
            book = load_workbook(path)
            self.assertEqual(book.sheetnames[0], cover); self.assertEqual(len(book.sheetnames), 8)
            self.assertEqual(book.worksheets[2].sheet_view.rightToLeft, language == "ar" or None)


if __name__ == "__main__":
    unittest.main()
