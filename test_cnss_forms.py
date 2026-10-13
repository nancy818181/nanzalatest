"""2.9.107: the 8 NSSF / MOF forms as editable Excel filled from Saber (replaces the PDF overlay tests)."""
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

import cnss_contribution_forms as contributions
from cnss_forms import FORMS, FORM_BY_KEY, LEGACY_FORM_KEYS, build_workbook, form_data, summary_rows
from payroll_excel_forms import leave_code

COMPANY = {"company_name": "Test Company", "company_nssf": "605123", "company_mof": "987654", "company_address": "Beirut",
           "company_phone": "01-123456"}
EMPLOYEE = {"id": 7, "employee_number": "E7", "full_name": "Ahmad Khalil", "father_name": "Mohamad", "mother_name": "Fatima",
            "birth_date": "1990-05-12", "birth_place": "Tripoli", "nationality": "Lebanese", "national_id": "000123", "nssf_number": "1234567",
            "marital_status": "married", "children": 2, "hire_date": "2026-01-02", "leave_date": "2026-09-30", "leave_reason": "استقالة",
            "base_salary": "40000000", "currency": "LBP", "sex": "male", "pay_type": "monthly", "addr_governorate": "North", "addr_street": "Main"}


def line(nssf, name, month, sick, family, eos, rate_sick=0.11, rate_family=0.06, rate_eos=0.085):
    employee_share, employer_sick = round(sick * 0.03), round(sick * (rate_sick - 0.03))
    fam, end = round(family * rate_family), round(eos * rate_eos)
    total = employee_share + employer_sick + fam + end
    return [nssf, name, month, sick, sick, employee_share, employer_sick, family, fam, eos, end, total, 0, total]


def report(rows, allowance=0, label="September 2026", start="2026-09-01", end="2026-09-30"):
    keys = ("salary", "sick_base", "employee", "employer_sick", "family_base", "family", "eos_base", "eos", "total", "allowance", "net")
    summary = {key: sum(row[3 + i] for row in rows) for i, key in enumerate(keys)}
    summary["allowance"] = allowance; summary["net"] = summary["total"] - allowance
    return {"report": "NSSF", "period_label": label, "date_from": start, "date_to": end, "record_count": len(rows),
            "payroll_employee_count": len({r[0] for r in rows}), "summary": summary,
            "monthly_detail": {"rows": rows + [["TOTAL | المجموع", "", ""] + [0] * 11]},
            "sections": [{}, {}, {"rows": [["01-2026 to 12-2026", "-", "-", "3%", "8%", "6%", "8.5%", "-"]]}, {}]}


class FormsCatalogueTest(unittest.TestCase):
    def test_eight_forms_and_old_keys_still_open_them(self):
        self.assertEqual(len(FORMS), 8)
        for old, new in LEGACY_FORM_KEYS.items():
            self.assertIn(new, FORM_BY_KEY, old)

    def test_every_form_builds_an_rtl_a4_workbook(self):
        folder = Path(tempfile.mkdtemp())
        sept = report([line("1234567", "Ahmad", "09-2026", 40_000_000, 40_000_000, 40_000_000)])
        for form in FORMS:
            with self.subTest(form=form.key):
                data, _notes = form_data(form.key, COMPANY, EMPLOYEE, sept, [], year=2026, period_type="monthly", index=9)
                path = folder / f"{form.key}.xlsx"
                build_workbook(form.key, data, path)
                book = load_workbook(path)
                sheet = book.worksheets[0]
                self.assertTrue(sheet.sheet_view.rightToLeft)
                self.assertEqual(int(sheet.page_setup.paperSize), int(sheet.PAPERSIZE_A4))
                self.assertTrue(summary_rows(form.key, data))


class Contributions190ATest(unittest.TestCase):
    def setUp(self):
        self.rows = [line("1234567", "Ahmad", "09-2026", 40_000_000, 40_000_000, 40_000_000),
                     line("7654321", "Rim", "09-2026", 30_000_000, 0, 30_000_000)]
        self.report = report(self.rows, allowance=1_500_000)

    def test_wages_times_rate_gives_the_payroll_contributions(self):
        data, _ = form_data("cnss_contributions_monthly_quarterly", COMPANY, report=self.report, year=2026, period_type="monthly", index=9)
        summary = self.report["summary"]
        expected = {"sick": summary["employee"] + summary["employer_sick"], "eos": summary["eos"], "family": summary["family"]}
        for branch, due in expected.items():
            with self.subTest(branch=branch):
                self.assertEqual(round(data[branch]["wages"] * data[branch]["rate"]), due)
        self.assertEqual(data["sick"]["count"], 2)
        self.assertEqual(data["family"]["count"], 1)  # Rim has no family-allowance base
        self.assertEqual(data["allowance"], 1_500_000)
        self.assertIn("أيلول", data["period_text"])

    def test_totals_are_formulas_and_amounts_are_numbers(self):
        data, _ = form_data("cnss_contributions_monthly_quarterly", COMPANY, report=self.report, year=2026, period_type="monthly", index=9)
        sheet = build_workbook("cnss_contributions_monthly_quarterly", data).worksheets[0]
        formulas = [c.value for row in sheet.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
        self.assertTrue(any("ROUND(E" in f and "*J" in f for f in formulas))       # contribution = wages x rate
        self.assertTrue(any(f.startswith("=SUM(L") for f in formulas))             # total contributions
        self.assertTrue(any(f.startswith("=L") and "-L" in f for f in formulas))   # net = total - allowances
        numbers = [c.value for row in sheet.iter_rows() for c in row if isinstance(c.value, (int, float))]
        self.assertIn(70_000_000, numbers)                                           # sickness wages as a number

    def test_payment_method_is_marked(self):
        data, notes = form_data("cnss_contributions_monthly_quarterly", COMPANY, report=self.report, year=2026, period_type="monthly", index=9,
                                payment_details={"method": "Bank transfer", "reference": "R-1"})
        values = [c.value for row in build_workbook("cnss_contributions_monthly_quarterly", data).worksheets[0].iter_rows() for c in row]
        self.assertIn("X", values)
        self.assertFalse(any("Payment method" in n for n in notes))

    def test_missing_payment_details_are_reported(self):
        _, notes = form_data("cnss_contributions_monthly_quarterly", COMPANY, report=self.report, year=2026, period_type="monthly", index=9)
        self.assertTrue(any("Payment method" in n for n in notes))
        self.assertTrue(any("receipt" in n for n in notes))


class AnnualFormsTest(unittest.TestCase):
    def setUp(self):
        rows = [line("1234567", "Ahmad", f"{m:02d}-2026", 40_000_000, 40_000_000, 40_000_000) for m in range(1, 10)]
        self.year = report(rows, label="2026", start="2026-01-01", end="2026-12-31")

    def test_386_uses_payroll_months_when_nothing_was_saved(self):
        data, notes = form_data("cnss_annual_declaration", COMPANY, report=self.year, filed_wages=[], year=2026)
        self.assertEqual(data["filed"][3]["sick"], 40_000_000)
        self.assertEqual(data["filed"][11]["sick"], 0)
        self.assertTrue(any("not saved" in n for n in notes))
        self.assertEqual(len(data["employees"]), 1)

    def test_386_section_b_refers_to_the_schedule_sheet(self):
        data, _ = form_data("cnss_annual_declaration", COMPANY, report=self.year, filed_wages=[], year=2026)
        book = build_workbook("cnss_annual_declaration", data)
        self.assertEqual(len(book.worksheets), 2)
        main, schedule = book.worksheets
        formulas = [c.value for row in main.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
        self.assertTrue(any(f.startswith(f"='{schedule.title}'!P") for f in formulas))
        self.assertFalse(any("__SCHEDULE_" in f for f in formulas))
        schedule_values = [c.value for row in schedule.iter_rows() for c in row]
        self.assertIn("1234567", schedule_values)

    def test_351_settlement_is_actual_minus_declared(self):
        filed = [{"month": m, "sickness_wages": 30_000_000, "family_wages": 40_000_000, "end_service_wages": 40_000_000} for m in range(1, 13)]
        data, _ = form_data("cnss_contributions_yearly", COMPANY, report=self.year, filed_wages=filed, year=2026)
        self.assertEqual(data["sick"]["wages"], 9 * 40_000_000 - 12 * 30_000_000)
        self.assertEqual(data["pay_until"], "31-03-2027")


class EmployeeFormsTest(unittest.TestCase):
    def test_employee_needed_and_leaving_reason_is_read(self):
        _, notes = form_data("cnss_r3", COMPANY, None)
        self.assertTrue(any("Select an employee" in n for n in notes))
        data, _ = form_data("cnss_termination_notice", COMPANY, EMPLOYEE)
        self.assertEqual(data["leave_code"], "1")
        self.assertEqual(data["addr_governorate"], "North")
        self.assertEqual(leave_code("Death"), "5")
        self.assertEqual(leave_code(""), "")

    def test_r3_1_needs_no_employee(self):
        data, notes = form_data("cnss_r3_employee_registration", COMPANY, None)
        self.assertFalse(any("Select an employee" in n for n in notes))
        build_workbook("cnss_r3_employee_registration", data)


if __name__ == "__main__":
    unittest.main()


class ArabicAndDatesTest(unittest.TestCase):
    """2.9.108: Arabic fields first, nationality translated, today's date next to the company stamp, leaving reasons."""
    def test_arabic_fields_are_used_and_nationality_translated(self):
        company = dict(COMPANY, company_name_ar="شركة الاختبار", company_cr_number="12345")
        employee = dict(EMPLOYEE, full_name_ar="بسام صابر", father_name_ar="فؤاد", nationality="lebanese")
        data, _ = form_data("cnss_employment_declaration", company, employee)
        self.assertEqual(data["company_name"], "شركة الاختبار")
        self.assertEqual((data["first_name"], data["last_name"]), ("بسام", "صابر"))
        self.assertEqual(data["father_name"], "فؤاد")
        self.assertEqual(data["mother_name"], "فاطمة")           # no Arabic given: written in Arabic (2.9.109)
        self.assertEqual(data["nationality"], "لبنانية")
        self.assertEqual(data["company_cr_number"], "12345")
        self.assertNotIn("foreign_nationality", data)

    def test_company_date_is_filled(self):
        from datetime import date
        data, _ = form_data("cnss_termination_notice", COMPANY, EMPLOYEE)
        values = [c.value for row in build_workbook("cnss_termination_notice", data).worksheets[0].iter_rows() for c in row]
        self.assertIn(date.today().strftime("%d-%m"), values)
        self.assertIn(date.today().strftime("%y"), values)

    def test_every_printed_leaving_reason_is_recognised(self):
        from payroll_excel_forms import LEAVE_REASONS
        self.assertEqual([leave_code(reason) for reason in LEAVE_REASONS], ["1", "2", "3", "4", "5", "6", "7"])


class ArabicFromEnglishTest(unittest.TestCase):
    """2.9.109: English employee details are written in Arabic on the forms; the spouse is on R3 and 2AA."""
    def test_english_details_come_out_in_arabic(self):
        from arabic_names import to_arabic
        self.assertEqual(to_arabic("bassam saber"), "بسام صابر")
        self.assertEqual(to_arabic("hazmieh", "place"), "الحازمية")
        self.assertEqual(to_arabic("sales manager", "job"), "مدير مبيعات")
        self.assertEqual(to_arabic("lebanese", "nationality"), "لبنانية")
        self.assertEqual(to_arabic("بسام"), "بسام")                         # Arabic stays as typed
        self.assertTrue(to_arabic("Zyxwq"))                                  # an unknown name is still written in Arabic letters

    def test_corrected_arabic_wins_and_spouse_is_filled(self):
        employee = dict(EMPLOYEE, full_name="Ahmad Khalil", full_name_ar="أحمد الخليل", father_name="Mohamad",
                        spouse_name="Rima", spouse_maiden_name="Haddad", spouse_birth_date="1990-02-03", spouse_birth_place="beirut",
                        spouse_nationality="lebanese", spouse_id="777")
        data, _ = form_data("cnss_r3", COMPANY, employee)
        self.assertEqual(data["first_name"], "أحمد")                        # the corrected Arabic is used
        self.assertEqual(data["father_name"], "محمد")                        # no Arabic typed: converted
        self.assertEqual((data["spouse_name"], data["spouse_maiden_name"]), ("ريما", "حداد"))
        values = [c.value for row in build_workbook("cnss_r3", data).worksheets[0].iter_rows() for c in row]
        for expected in ("ريما", "حداد", "لبنانية", "بيروت", "777"):
            self.assertIn(expected, values)
        family = build_workbook("cnss_employment_declaration", form_data("cnss_employment_declaration", COMPANY, employee)[0]).worksheets[1]
        self.assertIn("ريما حداد", [c.value for row in family.iter_rows() for c in row])
