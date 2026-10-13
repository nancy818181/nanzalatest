"""Official Lebanese payroll forms (MOF R3, R3-1; CNSS 2AA, 41A, leave notice) as editable Excel sheets,
filled from the company settings and the selected employee. Same layout as the printed forms (RTL, A4, one page).
Light-yellow cells are editable; options are marked with X."""
import re
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from datetime import date
from openpyxl.worksheet.page import PageMargins

INPUT = "FFFDE8"
thin = Side(style="thin", color="000000")
right = Alignment(horizontal="right", vertical="center", wrap_text=True, readingOrder=2)
center = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2)


class Form:
    def __init__(self, title, cols=24, width=4.6, color="000000", frame="000000", data=None):
        self.data = data or {}
        self.wb = Workbook(); self.ws = self.wb.active; self.ws.title = title
        self.ws.sheet_view.rightToLeft = True; self.ws.sheet_view.showGridLines = False
        self.N = cols; self.color = color; self.frame_color = frame
        from openpyxl.utils import get_column_letter as L
        self.L = L
        for c in range(1, cols + 1): self.ws.column_dimensions[L(c)].width = width

    def font(self, size=10, bold=False, under=False, color=None):
        return Font(name="Arial", size=size, bold=bold, underline="single" if under else None, color=color or self.color)

    def merge(self, r, c1, c2, r2=None):
        r2 = r2 or r
        if c1 != c2 or r != r2: self.ws.merge_cells(start_row=r, start_column=c1, end_row=r2, end_column=c2)
        return self.ws.cell(r, c1)

    def label(self, r, c1, c2, text, size=10, bold=False, under=False, align=right, r2=None):
        cell = self.merge(r, c1, c2, r2); cell.value = text; cell.font = self.font(size, bold, under); cell.alignment = align
        return cell

    def value(self, key):
        value = self.data.get(key) if key else None
        return "" if value in (None, "") else str(value)

    def field(self, r, c1, c2, r2=None, box=False, key=None):
        cell = self.merge(r, c1, c2, r2)
        if key: cell.value = self.value(key) or None
        cell.font = self.font(10, color="1F3B73")
        # 2.9.108: a one-line field shrinks a long value (a salary, a name) to fit its box instead of breaking it on two lines
        cell.alignment = right if (r2 or r) != r else Alignment(horizontal="right", vertical="center", shrink_to_fit=True, readingOrder=2)
        for rr in range(r, (r2 or r) + 1):
            for c in range(c1, c2 + 1):
                x = self.ws.cell(rr, c); x.fill = PatternFill("solid", fgColor=INPUT)
                if box:
                    x.border = Border(left=thin if c == c1 else None, right=thin if c == c2 else None,
                                      top=thin if rr == r else None, bottom=thin if rr == (r2 or r) else None)
                elif rr == (r2 or r):
                    x.border = Border(bottom=Side(style="dotted", color="000000"))
        return cell

    def boxes(self, r, c1, count, step=1, key=None):
        # Numbers read left to right even on Arabic forms: in a right-to-left sheet the highest column is on the
        # visual left, so the first digit goes there.
        digits = [ch for ch in self.value(key) if ch.isalnum()][:count] if key else []
        digits = list(reversed(digits + [""] * (count - len(digits))))
        for i in range(count):
            x = self.ws.cell(r, c1 + i * step); x.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            x.fill = PatternFill("solid", fgColor=INPUT); x.alignment = center; x.font = self.font(10, color="1F3B73")
            if digits[i]: x.value = digits[i]

    def option(self, r, c, text, number, text_cols=2, key=None, match=None):
        """Printed choice: the word, then a box holding its number (type X over the number to choose it).
        Returns the next free column."""
        self.label(r, c, c + text_cols - 1, text, 9, bold=True)
        x = self.ws.cell(r, c + text_cols); x.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        x.fill = PatternFill("solid", fgColor=INPUT); x.alignment = center; x.font = self.font(9, True, color="1F3B73")
        x.value = "X" if key and match is not None and self.value(key) == str(match) else number
        return c + text_cols + 1

    def checkbox(self, r, c, text, text_cols=2, box_first=True, key=None, match=None):
        """☐ choice without a number (R3 style)."""
        if box_first:
            x = self.ws.cell(r, c); self.label(r, c + 1, c + text_cols, text, 9)
        else:
            self.label(r, c, c + text_cols - 1, text, 9); x = self.ws.cell(r, c + text_cols)
        x.border = Border(left=thin, right=thin, top=thin, bottom=thin); x.fill = PatternFill("solid", fgColor=INPUT)
        x.alignment = center; x.font = self.font(10, True, color="1F3B73")
        if key and match is not None and self.value(key) == str(match): x.value = "X"
        return c + text_cols + 1

    def new_sheet(self, title):
        ws = self.wb.create_sheet(title); ws.sheet_view.rightToLeft = True; ws.sheet_view.showGridLines = False
        for c in range(1, self.N + 1): ws.column_dimensions[self.L(c)].width = self.ws.column_dimensions["A"].width
        self.ws = ws
        return ws

    def date(self, r, c1, captions=True, size=7, key=None):
        parts = re.findall(r"\d+", self.value(key)) if key else []
        if len(parts) == 3 and len(parts[0]) == 4: parts = parts[::-1]
        for i, name in enumerate(("اليوم", "الشهر", "السنة")):
            c = c1 + i * 3; cell = self.field(r, c, c + 1)
            if len(parts) == 3: cell.value = parts[i]
            if i < 2: self.label(r, c + 2, c + 2, "/", 10, align=center)
            if captions: self.label(r + 1, c, c + 1, name, size, align=center)

    def hline(self, r, c1=1, c2=None, style="thin"):
        for c in range(c1, (c2 or self.N) + 1):
            x = self.ws.cell(r, c); b = x.border
            x.border = Border(left=b.left, right=b.right, top=Side(style=style, color=self.frame_color), bottom=b.bottom)

    def vline(self, c, r1, r2):
        """Vertical line between column c and c+1 (in a right-to-left sheet Excel draws "right" on the visual left)."""
        for r in range(r1, r2 + 1):
            x = self.ws.cell(r, c); b = x.border
            x.border = Border(right=Side(style="thin", color=self.frame_color), left=b.left, top=b.top, bottom=b.bottom)

    def fill(self, r1, r2, c1, c2, color):
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1): self.ws.cell(r, c).fill = PatternFill("solid", fgColor=color)

    def frame(self, r1, r2, c1=1, c2=None, style="medium"):
        c2 = c2 or self.N; s = Side(style=style, color=self.frame_color)
        for r in range(r1, r2 + 1):
            for c in (c1, c2):
                x = self.ws.cell(r, c); b = x.border
                x.border = Border(left=s if c == c1 else b.left, right=s if c == c2 else b.right, top=b.top, bottom=b.bottom)
        for c in range(c1, c2 + 1):
            for r, top in ((r1, True), (r2, False)):
                x = self.ws.cell(r, c); b = x.border
                x.border = Border(left=b.left, right=b.right, top=s if top else b.top, bottom=s if not top else b.bottom)

    def heights(self, rows, value):
        for r in rows: self.ws.row_dimensions[r].height = value

    def page(self, last_row, last_col=None):
        ws = self.ws
        note = self.merge(last_row + 2, 1, self.N)
        note.value = "للتعبئة: الخانات الملوّنة بالأصفر الفاتح قابلة للكتابة (ضع X في المربع المناسب، ورقم واحد في كل مربع للأرقام). هذا السطر لا يُطبع."
        note.font = self.font(8, color="7F7F7F"); note.alignment = right
        ws.print_area = f"A1:{self.L(last_col or self.N)}{last_row}"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4; ws.page_setup.orientation = "portrait"
        ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.4, bottom=0.4); ws.print_options.horizontalCentered = True

    def finish(self, path, last_row, last_col=None):
        self.page(last_row, last_col)
        if path: self.wb.save(path)
        return self.wb


BLUE = "1F3A93"
BLUE = "1F3A93"

def head(f, title2, foot_right_title=True):
    f.label(1, 1, 8, "الصندوق الوطني", 15, True); f.label(2, 1, 8, "للضمان الإجتماعي", 15, True)
    f.label(4, 1, f.N, "إعـــــــــــلام", 14, True, align=center); f.label(5, 1, f.N, title2, 14, True, align=center)
    f.heights((1, 2), 22); f.heights((4, 5), 22)
    r = 7; f.label(r, 1, 8, "ان صاحب العمل الموقع أدناه ،", 9, True)
    r = 8; f.label(r, 1, 6, "المسؤول عن مؤسسة :", 9, True); f.field(r, 7, f.N, key="company_name")
    r = 9; f.label(r, 1, 11, "المسجلة في الصندوق الوطني للضمان الإجتماعي تحت رقم :", 9, True)
    f.field(r, 16, 23, box=True, key="company_nssf")
    r = 10; f.label(r, 1, 4, "سجل تجاري : مكان", 9, True); f.field(r, 5, 9, key="company_cr_place"); f.label(r, 10, 10, "رقم", 9, True); f.field(r, 11, 14, key="company_cr_number")
    f.label(r, 15, 16, "هاتف", 9, True); f.field(r, 17, f.N, key="company_phone")
    r = 11; f.label(r, 1, 4, "وعنوانها الكامل :", 9, True); f.field(r, 5, f.N, key="company_address")
    r = 12; f.field(r, 1, f.N)
    r = 13; f.label(r, 1, 12, "يصرح بأن الأجير المبينة هويته فيما يلي والمسجل في الصندوق تحت الرقم :", 9, True)
    f.field(r, 16, 23, box=True, key="nssf_number")
    r = 14; f.label(r, 1, 2, "الجنس :", 9, True); c = f.option(r, 3, "ذكر", "١", key="sex", match="male"); f.option(r, c + 1, "انثى", "٢", key="sex", match="female")
    for rr in range(7, 15): f.ws.row_dimensions[rr].height = 20

def family(f, r):
    f.label(r, 1, 3, "الوضع العائلي :", 9, True); c = 4
    for text, number, match in (("اعزب", "١", "single"), ("متأهل", "٢", "married"), ("أرمل", "٣", "widowed"), ("مطلق", "٤", "divorced"), ("هاجر", "٥", "separated")):
        c = f.option(r, c, text, number, key="marital_status", match=match) + 1

def office_boxes(f, r, dates_year=False):
    mid = f.N // 2
    f.hline(r); f.vline(mid, r, r + 4)
    f.label(r, 1, mid, "حقل مخصص للصندوق :", 9, True); f.label(r, mid + 1, f.N, "عولج في دائرة التسجيل :", 9, True)
    if not dates_year: f.field(r, mid + 7, f.N)
    f.label(r + 1, 1, 4, "عولج في مركز رقم :", 9, True); f.field(r + 1, 5, mid - 1)
    f.label(r + 2, 1, 2, "بتاريخ :", 9, True); f.field(r + 2, 3, mid - 1)
    f.label(r + 2, mid + 1, mid + 2, "بتاريخ :", 9, True); f.field(r + 2, mid + 3, f.N - 1)
    f.label(r + 3, 1, 5, "اسم وتوقيع المستخدم", 9, True); f.field(r + 3, 6, mid - 1)
    f.label(r + 3, mid + 1, mid + 5, "اسم وتوقيع المستخدم", 9, True); f.field(r + 3, mid + 6, f.N - 1)
    f.heights(range(r, r + 5), 20)
    return r + 5

def build_cnss_leave(data=None, path=None):
    f = Form("ترك أجير", color=BLUE, frame=BLUE, data=data)
    head(f, "عن ترك اجير عمله في المؤسسة (١)")
    r = 15; f.label(r, 1, 3, "اسـم الاجير :", 9, True); f.field(r, 4, 13, key="first_name"); f.label(r, 14, 15, "الشهرة :", 9, True); f.field(r, 16, f.N, key="last_name")
    r = 16; f.label(r, 1, 2, "اسم الأب :", 9, True); f.field(r, 3, 11, key="father_name"); f.label(r, 12, 15, "اسم الأم وشهرتها :", 9, True); f.field(r, 16, f.N, key="mother_name")
    r = 17; f.label(r, 1, 5, "تاريخ ومحل ولادة الاجير :", 9, True); f.field(r, 6, 11, key="birth"); f.label(r, 12, 14, "رقم السجل :", 9, True); f.field(r, 15, 18, key="registry_number")
    f.label(r, 19, 20, "الجنسية :", 9, True); f.field(r, 21, f.N, key="nationality")
    family(f, 19)
    r = 20; f.label(r, 1, 4, "ترك العمل بها منذ :", 9, True); f.field(r, 5, 14, key="leave_date")
    r = 21; f.label(r, 1, 3, "سبب ترك العمل :", 9, True); c = 4
    for text, number, code in (("استقالة", "١", "1"), ("بلوغ السن", "٢", "2"), ("عجز", "٣", "3"), ("زواج", "٤", "4"), ("وفاة", "٥", "5"), ("هجرة", "٦", "6"), ("عمل آخر", "٧", "7")):
        c = f.option(r, c, text, number, text_cols=2 if len(text) > 4 else 1, key="leave_code", match=code)
    r = 22; f.label(r, 1, 3, "عمل الأجير الحالي", 9, True); f.field(r, 4, 20, key="job_title"); f.boxes(r, 22, 3)
    r = 23; f.label(r, 1, 7, "ان راتب الأجير بتاريخ ترك العمل هو :", 9, True); f.field(r, 8, f.N, key="salary_text")
    f.heights(range(15, 24), 20)
    r = 25; f.label(r, 1, 2, "بيروت في", 9, True); f.field(r, 3, 6, key="today_dm"); f.label(r, 7, 7, "٢٠", 9, True); f.field(r, 8, 8, key="today_yy")
    f.label(r, 11, 14, "خاتم المؤسسة", 9, True, align=center); f.label(r, 17, f.N, "اسم المسؤول عن المؤسسة وتوقيعه", 9, True, align=center)
    f.field(26, 17, f.N, 29, box=True); f.field(26, 11, 14, 29, box=True); f.heights(range(26, 30), 18)
    r = 31; f.hline(r)
    f.label(r, 1, 9, "اوافق على صحة المعلومات المدرجة اعلاه :", 9, True); f.label(r, 13, 15, "اسم الاجير :", 9, True); f.field(r, 16, f.N, key="full_name")
    r = 32; f.label(r, 1, 2, "التاريخ", 9, True); f.field(r, 3, 6, key="city"); f.label(r, 7, 7, "في", 9); f.field(r, 8, 9, key="today_dm"); f.label(r, 10, 10, "٢٠", 9)
    f.field(r, 11, 11, key="today_yy")
    f.label(r, 13, 15, "التوقيــــع :", 9, True); f.field(r, 16, f.N); f.heights((31, 32), 22)
    r = office_boxes(f, 34)
    f.hline(r)
    f.label(r, 1, f.N, "(١) تملأ المؤسسة هذه المطبوعة، وترسلها الى الصندوق الوطني للضمان الإجتماعي، كلما ترك احد الاجراء المرتبطين بها عمله فيها، مهما كان السبب. وذلك ضمن مهلة (١٥) يوماً من تاريخ الترك (المادة ٨٠ فقرة \"٤\")",
            7, r2=r + 1)
    f.label(r + 2, 1, f.N, "* تنبيـــــه : تحت طائلة اهمال المطبوعة: يجب دائماً ذكر رقم المؤسسة ورقم الاجير أو تاريخ ولادته بخط واضح ومقروء.", 7)
    LAST = r + 2; f.frame(1, LAST)
    return f.finish(path, LAST)


def build_cnss_41a(data=None, path=None):
    f = Form("إعلام باستخدام أجير", color=BLUE, frame=BLUE, data=data)
    head(f, "عن استخدام أجير")
    r = 15; f.label(r, 1, 2, "الاسـم :", 9, True); f.field(r, 3, 8, key="first_name"); f.label(r, 9, 10, "الشهرة :", 9, True); f.field(r, 11, 14, key="last_name")
    f.label(r, 15, 16, "رقمه في الصندوق", 9, True); f.field(r, 18, 23, box=True, key="nssf_number")
    r = 16; f.label(r, 1, 2, "اسم الأب :", 9, True); f.field(r, 3, 11, key="father_name"); f.label(r, 12, 15, "اسم الأم وشهرتها :", 9, True); f.field(r, 16, f.N, key="mother_name")
    r = 17; f.label(r, 1, 5, "تاريخ ومحل الولادة :", 9, True); f.field(r, 6, 10, key="birth"); f.label(r, 11, 13, "رقم السجل :", 9, True); f.field(r, 14, 16, key="registry_number")
    f.label(r, 17, 18, "الجنسية :", 9, True); f.field(r, 19, 21, key="nationality"); f.boxes(r, 23, 2)
    family(f, 19); f.boxes(19, 23, 2)
    r = 20; f.label(r, 1, 4, "استخدم فيها منذ :", 9, True); f.field(r, 5, 12, key="hire_date"); f.label(r, 13, 17, "عدد ساعات العمل في الشهر", 9, True); f.field(r, 18, f.N, key="work_hours")
    r = 21; c = f.option(r, 1, "بدوام كامل", "١", key="work_time", match="full"); f.option(r, c, "دوام جزئي", "٢", key="work_time", match="part")
    r = 22; f.label(r, 1, 4, "* عمل الأجير الحالي", 9, True); f.field(r, 5, 10, key="job_title"); f.boxes(r, 11, 3)
    f.label(r, 15, 17, "ان الراتب الحالي :", 9, True); f.field(r, 18, 22, key="salary_lbp"); f.label(r, 23, 23, "ل.ل.", 9, True)
    r = 23; f.label(r, 1, 3, "* طريقة دفع الأجر :", 9, True); c = 4
    for text, number, m in (("شهري", "١", "monthly"), ("اسبوعي", "٢", "weekly"), ("يومي", "٣", "daily"), ("لقاء عمولة", "٤", "commission"), ("على الانتاج", "٥", "piece")):
        c = f.option(r, c, text, number, key="pay_type", match=m)
    r = 24; f.label(r, 1, 9, "* هل يعمل حسب معرفتك لدى صاحب عمل آخر :", 9, True); c = f.option(r, 10, "نعم", "١", 1, key="other_job", match="yes")
    c = f.option(r, c, "كلا", "٢", 1, key="other_job", match="no")
    f.label(r, c, f.N, "اذكر اسم المؤسسة ورقمها", 9, True)
    r = 25; f.field(r, 1, 14, key="other_employer"); f.field(r, 16, 19, box=True, key="other_employer_nssf"); f.field(r, 20, 21, box=True); f.field(r, 22, 23, box=True)
    r = 26; f.label(r, 1, 11, "* اسم وعنوان صاحب العمل الذي كان الاجير يعمل لديه قبل دخوله مؤسستي :", 9, True); f.field(r, 12, f.N, key="previous_employer")
    r = 27; f.label(r, 1, 3, "رقمه في الصندوق", 9, True); f.field(r, 4, 14, key="previous_employer_nssf"); f.field(r, 16, 19, box=True); f.field(r, 20, 21, box=True); f.field(r, 22, 23, box=True)
    f.heights(range(15, 28), 20)
    r = 29; f.hline(r); f.label(r, 1, 9, "* اوافق على صحة المعلومات المدرجة اعلاه :", 9, True)
    f.label(r, 11, 14, "ختم المؤسسة", 9, True, align=center); f.label(r, 16, f.N, "اسم وتوقيع الشخص المسؤول", 9, True, align=center)
    r = 30; f.label(r, 1, 3, "اسم الاجير", 9, True); f.field(r, 4, 6, key="full_name"); f.label(r, 7, 8, "التوقيع", 9, True); f.field(r, 9, 10)
    f.field(30, 11, 14, 32, box=True); f.field(30, 16, f.N, 32, box=True)
    r = 33; f.field(r, 1, 4, key="city"); f.label(r, 5, 5, "في", 9); f.field(r, 6, 7, key="today_dm"); f.label(r, 8, 8, "٢٠", 9); f.field(r, 9, 9, key="today_yy")
    f.field(r, 14, 17, key="city"); f.label(r, 18, 18, "في", 9); f.field(r, 19, 20, key="today_dm"); f.label(r, 21, 21, "٢٠", 9); f.field(r, 22, 22, key="today_yy")
    f.heights(range(29, 34), 20)
    r = office_boxes(f, 35)
    f.hline(r)
    f.label(r, 1, f.N, "* تملأ هذه المطبوعة وترسل الى الصندوق الوطني للضمان الإجتماعي كلما استخدمت المؤسسة اجيراً جديداً سبق ان جرى تسجيله في الصندوق وذلك ضمن مهلة ١٥ يوماً من تاريخ الاستخدام",
            7, r2=r + 1)
    f.label(r + 2, 1, f.N, "* تنبيه : يهمل كل اعلام غير مرفق باخراج قيد فردي للعازب وعائلي للمتأهل، وغير مكتوب بخط واضح ومقروء، واذا كانت الكلمات ناقصة.", 7)
    f.label(r + 3, 19, f.N, "CNSS 41A", 9, True)
    LAST = r + 3; f.frame(1, LAST)
    return f.finish(path, LAST)

def build_cnss_2aa(data=None, path=None):
    f = Form("تصريح ص1", data=data)
    N = f.N
    f.label(1, 1, N, "تصريح باستخدام اجير", 16, True, align=center)
    f.label(2, 1, N, "(يملأ هذا التصريح من قبل صاحب العمل وعلى مسؤوليته)", 11, align=center); f.heights((1, 2), 24)
    P = 17          # main part: columns 1..P (number column 1); right panel (visual left): P+1..N
    top = 4; f.hline(top, style="medium")
    def num(r, n): f.label(r, 1, 1, n, 9, align=center)
    r = 5; num(r, "١"); f.label(r, 2, 9, "أ – معلومات خاصة بالمؤسسة:", 9, True); f.field(r, 12, 16, box=True, key="company_nssf")
    r = 6; f.label(r, 2, 11, "اسم وشهرة صاحب العمل أو صاحب العمل أو اسم الشركة", 9); f.field(r, 12, P, key="company_name")
    r = 7; f.field(r, 2, P)
    r = 8; f.label(r, 2, 4, "سجل تجاري: مكان", 9); f.field(r, 5, 8, key="company_cr_place")
    f.label(r, 9, 9, "رقم", 9); f.field(r, 10, 12, key="company_cr_number"); f.label(r, 13, 13, "هاتف", 9); f.field(r, 14, P, key="company_phone")
    r = 9; num(r, "٢"); f.label(r, 2, 8, "عنوان المؤسسة (مكان عمل الاجير)", 9); f.field(r, 9, P, key="company_address")
    r = 10; f.field(r, 2, P)
    f.hline(11, 1, P)
    r = 12; f.label(r, 2, P, "ب – معلومات خاصة بالاجير (ضع علامة X في المربع المناسب):", 9, True)
    r = 13; num(r, "٣"); f.label(r, 2, 3, "الجنس:", 9); c = f.option(r, 4, "ذكر", "١", 1, key="sex", match="male"); f.option(r, c + 1, "انثى", "٢", 1, key="sex", match="female")
    r = 14; num(r, "٤"); f.label(r, 2, 3, "اسم الاجير:", 9); f.field(r, 4, 9, key="first_name"); f.label(r, 10, 11, "الشهرة", 9); f.field(r, 12, P, key="last_name")
    r = 15; num(r, "٥"); f.label(r, 2, 3, "اسم الأب:", 9); f.field(r, 4, 9, key="father_name"); f.label(r, 10, 12, "اسم الام وشهرتها", 9); f.field(r, 13, P, key="mother_name")
    r = 16; num(r, "٦"); f.label(r, 2, 5, "تاريخ ومحل الولادة:", 9); f.field(r, 6, 9, key="birth"); f.label(r, 10, 11, "القضاء", 9); f.field(r, 12,
            13, key="birth_caza"); f.label(r, 14, 15, "رقم السجل:", 9); f.field(r, 16, P, key="registry_number")
    r = 17; num(r, "٧"); f.label(r, 2, 3, "الوضع العائلي:", 9); c = 4
    for text, n, match in (("أعزب", "١", "single"), ("متأهل", "٢", "married"), ("ارمل", "٣", "widowed"), ("مطلق", "٤", "divorced"), ("هاجر", "٥", "separated")):
        c = f.option(r, c, text, n, 1, key="marital_status", match=match)
    r = 18; num(r, "٨"); f.label(r, 2, 5, "محل الاقامة (حسب الهوية)", 9); f.field(r, 6, 11, key="registry_place"); f.label(r, 12, 13, "المذهب", 9); f.field(r, 14, P, key="sect")
    r = 19; num(r, "٩"); f.label(r, 2, 6, "عنوان السكن الحالي: المحافظة", 9); f.field(r, 7, 11, key="addr_governorate"); f.label(r, 12, 13, "القضاء", 9); f.field(r, 14, P, key="addr_caza")
    r = 20; f.label(r, 2, 4, "المدينة أو القرية:", 9); f.field(r, 5, 8, key="addr_town"); f.label(r, 9, 9, "الحي", 9); f.field(r, 10, 12, key="addr_district"); f.label(r, 13, 14, "الشارع", 9)
    f.field(r, 15, P, key="addr_street")
    r = 21; f.label(r, 2, 4, "بناية وطابق", 9); f.field(r, 5, 11, key="building_floor"); f.label(r, 12, 13, "هاتف", 9); f.field(r, 14, P, key="contact_number")
    r = 22; num(r, "١٠"); f.label(r, 2, 4, "تاريخ دخول العمل", 9); f.field(r, 5, 9, key="hire_date"); f.label(r, 10, 13, "عدد ساعات عمله في الشهر", 9); f.field(r, 14, P, key="work_hours")
    r = 23; f.label(r, 2, 3, "دوام العمل:", 9); c = f.option(r, 4, "كامل", "١", 1, key="work_time", match="full"); f.option(r, c + 1, "جزئي", "٢", 1, key="work_time", match="part")
    r = 24; num(r, "١١"); f.label(r, 2, 4, "عمل الاجير الحالي", 9); f.field(r, 5, 8, key="job_title"); f.boxes(r, 9, 3); f.label(r, 12, 13, "الراتب الحالي:",
            9); f.field(r, 14, 16, key="salary_lbp"); f.label(r, P, P, "ل.ل.", 9)
    r = 25; num(r, "١٢"); f.label(r, 2, 5, "الاجر بتاريخ دخول العمل:", 9); f.field(r, 6, 10, key="hire_salary"); f.label(r, 11, 11, "ل.ل.", 9)
    r = 26; f.label(r, 2, 3, "طريقة دفع الاجر:", 9); c = 4
    for text, n, m in (("شهري", "١", "monthly"), ("اسبوعي", "٢", "weekly"), ("يومي", "٣", "daily"), ("لقاء عمولة", "٤", "commission"), ("على الانتاج", "٥", "piece")):
        c = f.option(r, c, text, n, 1 if len(text) < 6 else 2, key="pay_type", match=m)
    r = 27; num(r, "١٣"); f.label(r, 2, 6, "اذا كان الاجير اجنبياً تذكر", 9); f.label(r, 7, 8, "الجنسية", 9); f.field(r, 9, P, key="foreign_nationality")
    r = 28; f.label(r, 2, 6, "رقم وتاريخ اجازة عمله", 9); f.field(r, 7, P, key="work_permit")
    r = 29; num(r, "١٤"); f.label(r, 2, 10, "هل يعمل الاجير، حسب معرفتك، لدى صاحب عمل آخر؟", 9); c = f.option(r, 11, "نعم", "١", 1, key="other_job", match="yes")
    f.option(r, c, "كلا", "٢", 1, key="other_job", match="no")
    r = 30; num(r, "١٥"); f.label(r, 2, 6, "اسم ورقم صاحب العمل الآخر:", 9); f.field(r, 7, 11, key="other_employer"); f.field(r, 12, 13, box=True, key="other_employer_nssf")
    f.field(r, 14, 14, box=True); f.field(r, 15, 16, box=True)
    f.heights(range(5, 31), 19)
    # right panel (left side of the page): for the Fund
    f.vline(P, 4, 30)
    f.label(5, P + 2, N, "اترك هذا الحقل فارغاً", 9, align=center)
    f.checkbox(8, P + 2, "الاجير مسجل سابقاً", 4); f.label(10, P + 2, N, "تحت الرقم", 9)
    f.field(11, P + 2, P + 4, box=True); f.field(11, P + 5, N - 1, box=True)
    f.checkbox(15, P + 2, "الاجير غير مسجل سابقاً في الصندوق", 5); f.ws.row_dimensions[15].height = 28
    f.label(16, P + 2, P + 4, "اسم المستخدم:", 9); f.field(16, P + 5, N)
    f.label(17, P + 2, P + 4, "التوقيـــع:", 9); f.field(17, P + 5, N)
    f.label(19, P + 2, N, "بعد التدقيق", 9); f.label(20, P + 2, N, "اعطي رقم التسجيل", 9)
    f.field(21, P + 2, P + 4, box=True); f.field(21, P + 5, N - 1, box=True)
    f.label(24, P + 2, P + 3, "بتاريخ:", 9); f.field(24, P + 4, N)
    f.label(25, P + 2, N, "اسم المستخدم وتوقيعه", 9); f.field(26, P + 2, N, 27, box=True)
    f.hline(31, style="medium")
    r = 32; f.label(r, 1, 9, "اوافق على صحة المعلومات المدرجة أعلاه: اسم الاجير:", 9); f.field(r, 10, 14, key="full_name"); f.label(r, 15, 16, "التاريخ",
            9); f.field(r, 17, 19, key="today"); f.label(r, 20, 21, "التوقيع", 9); f.field(r, 22, N)
    f.hline(33, style="thin")
    r = 34; f.label(r, 12, 15, "ختم المؤسسة", 9, align=center); f.label(r, 17, N, "اسم وتوقيع الشخص المسؤول", 9, align=center)
    f.field(35, 12, 15, 38, box=True); f.field(35, 17, N, 38, box=True)
    r = 39; f.field(r, 14, 18, key="city"); f.label(r, 19, 19, "في", 9); f.field(r, 20, N, key="today")
    f.hline(41)
    f.label(41, 1, N, "١ – ملاحظة: كل صاحب عمل يغفل التصريح عن اجرائه ضمن مهلة (١٥) يوماً من تاريخ الاستخدام أو يتقدم بتصاريح غير صحيحة يتعرض لعقوبات الغرامة والحبس وذلك عملاً باحكام المادتين ٨٠ و ٨١ من قانون الضمان الاجتماعي.",
            7, r2=42)
    f.label(43, 1, N, "٢ – يجب أن تطابق المعلومات الواردة أعلاه مع اخراج القيد المرفق (افرادي للأعزب – عائلي للمتأهل).", 7)
    f.label(44, 1, N - 4, "٣ – تنبيه: يهمل كل طلب غير مكتوب بخط واضح ومقروء ومعلومات ناقصة.", 7); f.label(44, N - 3, N, "CNSS-2 AA", 8, True)
    f.page(44)

    # ===================== page 2: family members
    f.new_sheet("تصريح ص2 أفراد العائلة")
    f.label(1, 1, 18, "تصريح عن افراد عائلة الاجير او الاجيرة للإستفادة حالياً", 13, True, align=center)
    f.label(2, 1, 18, "من التعويضات العائلية وفي المستقبل من ضمان المرض", 13, True, align=center)
    f.label(3, 1, 18, "(يملأ هذا التصريح من قبل الاجير وعلى مسؤوليته)", 10, align=center)
    f.label(1, 20, N, "مخصص للصندوق", 9, align=center); f.field(2, 20, 21, box=True); f.field(2, 22, 22, box=True); f.field(2, 23, N, box=True)
    f.heights((1, 2, 3), 22)
    intro = ("أن الأجير المدون اسمه على ظهر هذه الصفحة يصرح بأن أفراد عائلته الذين يذكر أسماؤهم في العامود الثالث من الجدول التالي تتوفر فيهم "
             "الشروط المطلوبة للاستفادة من التعويضات العائلية ومن تعويض المريض اي:\nأ – انهم يعيشون معه تحت سقف واحد وعلى نفقته لانهم لا يمارسون اي عمل مأجور وليس لديهم أي مدخول شخصي.\n"
             "ب – وتتوفر في كل منهم الشروط الخاصة المذكورة في العامود الثاني من الجدول نفسه.")
    f.label(5, 1, N, intro, 9, r2=8)
    # table: columns (visual right→left) who | conditions | name | sex | date & place of birth
    W = [(1, 3), (4, 10), (11, 16), (17, 17), (18, N)]
    heads = ("الافراد الذين يمكن ان يستفيدوا من الضمان", "الشروط الخاصة بكل واحد من افراد العائلة", "أسم وشهرة افراد العائلة الذي تتوفر فيهم جميع الشروط", "ذكر أو انثى", "تاريخ ومحل الولادة")
    r = 10
    for (c1, c2), text in zip(W, heads): f.label(r, c1, c2, text, 8, True, align=center); 
    f.ws.row_dimensions[r].height = 44
    from openpyxl.styles import Border, Side
    thin = Side(style="thin")
    def grid(r1, r2):
        for rr in range(r1, r2 + 1):
            for c1, c2 in W:
                for c in range(c1, c2 + 1):
                    x = f.ws.cell(rr, c); b = x.border
                    x.border = Border(top=thin if rr == r1 or True else b.top, bottom=thin if rr == r2 else b.bottom,
                                      left=thin if c == c1 else b.left, right=thin if c == c2 else b.right)
    grid(10, 10)
    rows = [("والد الاجير\nوالدة الاجير", "ان يكون عمره اكثر من ستين سنة\nاو ان يكون عاجزاً", ["ذكر", "انثى"]),
            ("زوجة الاجير", "لا يوجد شرط خاص. أما في حال تعدد الزوجات تذكر الزوجة الأولى فقط", ["انثى"]),
            ("زوج الاجيرة", "ان يكون عمره اكثر من ستين سنة\nأو ان يكون عاجزاً", ["ذكر"])]
    r = 11
    for who, cond, sexes in rows:
        n = len(sexes); f.label(r, 1, 3, who, 8, align=center, r2=r + n - 1); f.label(r, 4, 10, cond, 8, align=center, r2=r + n - 1)
        spouse_row = (who == "زوجة الاجير" and f.value("sex") != "female") or (who == "زوج الاجيرة" and f.value("sex") == "female")  # 2.9.109
        for i, sex in enumerate(sexes):
            f.field(r + i, 11, 16, key="spouse_full" if spouse_row else None); f.label(r + i, 17, 17, sex, 8, align=center)
            f.field(r + i, 18, N, key="spouse_birth" if spouse_row else None)
            f.ws.row_dimensions[r + i].height = 24 if n > 1 else 34
        grid(r, r + n - 1)
        for i in range(n): grid(r + i, r + i)
        r += n
    children = r
    f.label(r, 1, 3, "اولاد الاجير الشرعيون او بالتبني", 9, True, align=Alignment(horizontal="center", vertical="center", text_rotation=90, wrap_text=True), r2=r + 10)
    f.label(r, 4, 10, "– اولاد لم يبلغوا ١٦ سنة مكتملة:\nاذكر الاولاد من هذه الفئة ما عدا العاجزين منهم\n\n– اولاد عمرهم بين ١٦ و ٢٥ سنة مكتملة:\n"
            "اذكر الاناث والذكور الذين ينصرفون للدراسة والعازبات اللواتي لا يتابعن دروسهن ما عدا العجزة منهن (ضع علامة X امام اسماء الطلاب)\n\n– اولاد عجزة مهما كان عمرهم وجنسهم", 8, r2=r + 10)
    for i in range(11):
        f.field(r + i, 11, 16); f.field(r + i, 17, 17); f.field(r + i, 18, N); f.ws.row_dimensions[r + i].height = 21
        grid(r + i, r + i)
    grid(children, children + 10)
    r = children + 12
    f.label(r, 1, N, "ملاحظة: ترفق بهذا التصريح شهادة طبية تثبت حالة العجز، وإفادة مدرسية للطلاب الذين يتراوح عمرهم بين ١٦ و ٢٥ سنة", 8, r2=r)
    f.label(r + 1, 1, N, "حـــالات خـــاصـــة: – اذا كان هذا التصريح مقدم من قبل شخص غير والد او والدة الاولاد المصرح عنهم عليه أن يبين صفته بالنسبة للأولاد (التبني، الوصاية..) مع تقديم الأوراق الثبوتية",
            8, r2=r + 2); f.field(r + 3, 1, N)
    r += 4
    f.label(r, 1, N, "– اذا كان هذا التصريح مقدم من قبل أجيرة عليها أن تبين اذا كانت:", 8)
    c = 1; f.label(r + 1, 1, 6, "أ – (ضع علامة X في المربع المناسب)", 8)
    c = 7
    for text in ("عازبة", "متأهلة", "أرملة", "هاجرة", "مطلقة"):
        c = f.checkbox(r + 1, c, text, 2, box_first=False)
    f.label(r + 2, 1, N, "ب – اذا عهدت اليها رعاية الاولاد عليها أن ترفق مع تصريحها نسخة مصدقة عن الحكم الذي قضى بذلك.", 8)
    f.label(r + 3, 1, 8, "ج – هل زوج الاجيرة هو أيضاً اجير؟", 8); c = f.checkbox(r + 3, 9, "نعم", 1, box_first=False); f.checkbox(r + 3, c, "كلا", 1, box_first=False)
    f.label(r + 4, 1, 10, "في حالة النفي الرجاء تحديد الوضع المهني للزوج بالتفصيل", 8); f.field(r + 4, 11, N)
    f.label(r + 5, 1, N, "كل أجير يتعمد تقديم معلومات غير صحيحة يتعرض لعقوبات الحبس والغرامة وفقاً للمادتين ٨٠ و ٨١ من قانون الضمان.", 8, True)
    r += 7
    f.label(r, 1, 9, "مخصص للأحوال الشخصية", 9, align=center); f.field(r + 1, 1, 9, r + 4, box=True)
    f.field(r, 13, 18); f.label(r, 19, 19, "في", 9); f.field(r, 20, N); f.label(r + 3, 16, N, "التوقيع", 9, align=center); f.field(r + 4, 16, N)
    return f.finish(path, r + 4)

def build_mof_r3(data=None, path=None):
    GREEN = "5DBF5D"; GL = "3A9A3A"
    f = Form("ر3", frame=GL, data=data); N = f.N
    f.fill(1, 4, 1, N, GREEN)
    f.label(1, 1, 7, "الجمهورية اللبنانيــــة", 10, True); f.label(2, 1, 7, "وزارة الماليـــــة", 10, True)
    f.label(3, 1, 7, "مديرية المالية العامة", 8); f.label(4, 1, 10, "مديرية الواردات – ضريبة الرواتب والأجور", 8)
    c = f.label(1, 8, 20, "طلب تســـجيل مســـتخدم/أجيـــر جـــديد", 18, True, align=center, r2=3); c.fill = PatternFill("solid", fgColor=GREEN)
    c = f.label(1, 21, N, "ر ٣", 18, True, align=center, r2=3); c.fill = PatternFill("solid", fgColor=GREEN)
    f.heights(range(1, 5), 15)
    r = 6; f.label(r, 1, 5, "إســـــم الشركة/المؤسسة", 10, True, True); f.field(r, 6, 13, key="company_name"); f.label(r, 14, 16, "الشهرة التجارية", 9); f.field(r, 17, N, key="trade_name")
    r = 7; f.label(r, 1, 5, "رقم تسجيل الشركة/المؤسسة", 10, True); f.boxes(r, 6, 10, key="company_mof"); f.label(r + 1, 1, 5, "(لدى وزارة المالية)", 8)
    f.hline(9)
    r = 10; f.label(r, 1, 5, "تعـــريف المستخدم / الأجير", 10, True, True); f.label(r, 6, 11, "هل لديه رقم مالي شخصي؟*", 9)
    c = f.checkbox(r, 12, "نعم", 1, key="has_mof", match="yes"); c = f.checkbox(r, c, "كلا", 1, key="has_mof", match="no"); f.label(r, c, c + 3,
            "في حال نعم، أذكر الرقم", 8); f.boxes(r, c + 4, N - c - 3, key="mof_number")
    f.label(r + 1, 6, 11, "(لدى وزارة المالية)", 8)
    def item(r, c1, text, key, f2, lw=3):
        f.label(r, c1, c1 + lw - 1, text, 9); f.field(r, c1 + lw, f2, key=key or None)
    r = 12; item(r, 1, "١. الإســـم", "first_name", 11); item(r, 13, "٢. الشـــهرة", "last_name", N)
    r = 13; item(r, 1, "٣. إسم الأب", "father_name", 11); item(r, 13, "٤. إسم الأم وشهرتها قبل الزواج", "mother_name", N, 5)
    r = 14; f.label(r, 1, 3, "٥. الجنـس*", 9); c = f.checkbox(r, 4, "ذكر", 1, key="sex", match="male"); f.checkbox(r, c, "أنثى", 1, key="sex",
            match="female"); item(r, 13, "٦. الجنسية", "nationality", N)
    r = 15; item(r, 1, "٧. محل الــولادة", "birth_place", 11); f.label(r, 13, 15, "٨. تاريخ الولادة", 9); f.date(r, 16, key="birth_date")
    r = 17; item(r, 1, "٩. رقم السجل", "registry_number", 6); item(r, 7, "١٠. مكان السجل", "registry_place", 11); item(r, 13, "١١. رقم بطاقة الهوية", "national_id", N, 4)
    r = 18; f.label(r, 1, 3, "١٢. الوضع العائلي*", 9); c = 4
    for t, m in (("أعزب", "single"), ("متزوج", "married"), ("أرمل", "widowed"), ("مطلَق", "divorced")): c = f.checkbox(r, c, t, 1, key="marital_status", match=m)
    item(r, 13, "١٣. عدد الأولاد", "children", 17, 3)
    r = 19; f.label(r, 1, 4, "١٤. تاريخ إبتداء العمل", 9); f.date(r, 5, key="hire_date"); item(r, 15, "١٥. رقم الضمان الإجتماعي", "nssf_number", N, 4)
    r = 21; f.label(r, 13, 15, "١٦. نوع الأجر*", 9); c = 16
    for t, m in (("شهري", "monthly"), ("يومي", "daily"), ("بالساعة", "hourly")): c = f.checkbox(r, c, t, 2 if t == "بالساعة" else 1, key="pay_type", match=m)
    f.heights(range(10, 22), 19)
    f.hline(22)
    r = 23; f.label(r, 1, 8, "معلومات خاصة بالزوج/الزوجة", 10, True, True)
    r = 24; item(r, 1, "١. إسم الزوج/الزوجة", "spouse_name", 11, 4); item(r, 13, "٢. الشهرة قبل الزواج", "spouse_maiden_name", N, 4)
    r = 25; item(r, 1, "٣. إسم الأب", "spouse_father", 11); item(r, 13, "٤. إسم الأم وشهرتها قبل الزواج", "spouse_mother", N, 5)
    r = 26; item(r, 1, "٥. الجنسية", "spouse_nationality", 11); item(r, 13, "٦. محل الــولادة", "spouse_birth_place", N)
    r = 27; f.label(r, 1, 3, "٧. تاريخ الولادة", 9); f.date(r, 4, key="spouse_birth_date"); item(r, 13, "٨. رقم بطاقة الهوية", "spouse_id", N, 4)
    r = 29; f.label(r, 1, 10, "٩. عدد الأشخاص الذين يستفيدون من التنزيل العائلي", 9); f.field(r, 11, 13, key="family_dependents")
    r = 30; f.label(r, 1, 5, "١٠. هل الزوج/الزوجة يعمل*؟", 9, True); c = f.checkbox(r, 6, "نعم", 1, key="spouse_works", match="yes"); c = f.checkbox(r, c, "لا",
            1, key="spouse_works", match="no"); f.label(r, c, c + 3, "في حال نعم أذكر:", 9)
    r = 31; f.label(r, 3, 6, "رقم التسجيل الشخصي", 9); f.boxes(r, 7, 10, key="spouse_mof"); f.label(r + 1, 3, 6, "(لدى وزارة المالية)", 8)
    r = 33; f.label(r, 1, 11, "أ – في القطاع الخاص أو في مؤسسة عامة/مصلحة مستقلة", 9, True, True); f.label(r, 15, 17, "رقم التسجيل", 9); f.boxes(r, 18, 7, key="spouse_employer_mof")
    r = 34; f.label(r, 2, 9, "إسم الشركة/المؤسسة العامة/المصلحة المستقلة", 8); f.field(r, 10, 17, key="spouse_employer")
    r = 35; f.label(r, 1, 5, "ب – في الإدارات العامة", 9, True, True); f.label(r, 6, 7, "إسم الإدارة", 8); f.field(r, 8, 17, key="spouse_public_admin")
    f.heights(range(23, 36), 19)
    f.hline(36)
    r = 37; f.label(r, 1, 4, "عنوان السكن**", 10, True, True)
    r = 38
    for c1, t, c2, k in ((1, "محافظة", 5, "addr_governorate"), (7, "قضـــاء", 11, "addr_caza"), (13, "منطقة – بلدة", 17, "addr_town"), (19, "الحـــي", N, "addr_district")):
        f.label(r, c1, c1 + 1, t, 9); f.field(r, c1 + 2, c2, key=k)
    r = 39; f.label(r, 1, 2, "الشـــارع", 9); f.field(r, 3, 11, key="addr_street"); f.label(r, 12, 14, "المنطقة العقارية", 9); f.field(r, 15, 18); f.label(r, 19, 21, "رقم العقار/القسم", 9)
    f.field(r, 22, N)
    r = 40; f.label(r, 1, 2, "المبنـــى", 9); f.field(r, 3, 8, key="addr_building"); f.label(r, 9, 10, "الطابـــق", 9); f.field(r, 11, 13, key="addr_floor")
    f.label(r, 14, 14, "هاتف", 9); f.field(r, 15, 18, key="contact_number"); f.label(r, 19, 19, "هاتف", 9); f.field(r, 20, N, key="phone2")
    r = 41; f.label(r, 1, 4, "صندوق البريد: رقم", 9); f.field(r, 5, 7); f.label(r, 8, 9, "المنطقة", 9); f.field(r, 10, 16); f.label(r, 19, 19, "فاكس", 9); f.field(r, 20, 20); f.field(r, 21, N)
    r = 42; f.label(r, 1, 6, "البريد الإلكتروني (e-mail):", 9); f.field(r, 7, 20)
    f.heights(range(37, 43), 19)
    # إفادة (right) | خاص بالإدارة (left)
    r = 44; mid = 12
    f.fill(r, r, 1, N, "D9D9D9"); f.hline(r); f.vline(mid, r, r + 8)
    f.label(r, 1, mid, "إفادة", 10, True, align=center); f.label(r, mid + 1, N, "خاص بالإدارة", 10, True, align=center)
    f.label(r + 1, 1, mid, "أنا الموقع أدناه أفيد بأن هذه المعلومات هي مطابقة لتلك المدونة في بيان المعلومات من المستخدم / الأجير إلى صاحب العمل.", 9, r2=r + 2)
    f.label(r + 3, 1, 3, "إسم صاحب العمل", 9); f.field(r + 3, 4, mid - 1, key="company_name")
    f.label(r + 4, 1, 3, "الصفة", 9); f.field(r + 4, 4, mid - 1)
    f.label(r + 5, 1, 3, "التوقيع", 9); f.field(r + 5, 4, mid - 1)
    f.label(r + 6, 1, 2, "التاريخ", 9); f.date(r + 6, 3, key="today")
    f.label(r + 1, mid + 1, N, "الرقم المالي الشخصي للمستخدم/الأجير (لدى وزارة المالية)", 9)
    f.boxes(r + 2, mid + 2, 10)
    f.label(r + 4, mid + 1, mid + 3, "تـــاريخ التسجيل", 9); f.date(r + 4, mid + 4)
    f.label(r + 6, mid + 1, mid + 4, "أو سبب رفض التسجيل", 9); f.field(r + 6, mid + 5, N)
    f.heights(range(r, r + 9), 20)
    LAST = r + 8; f.frame(1, LAST)
    notes = ["* توضع علامة × في المربع المناسب", "** في حال تغيير عنوان السكن، يرجى تعبئة نموذج تعديل معلومات أفراد –م٥.",
             "– يستعمل هذا النموذج لتسجيل المستخدمين والأجراء الذين ليس لديهم رقم مالي شخصي لدى وزارة المالية ويقدم لدى الدائرة المالية المختصة.",
             "– يرفق بهذا الطلب صورة عن اخراج قيد عائلي للأجير المتزوج أو صورة عن بطاقة الهوية واخراج قيد افرادي للعازب."]
    for i, t in enumerate(notes): f.label(LAST + 1 + i, 1, N, t, 7)
    LAST += len(notes)
    f.ws.column_dimensions["Y"].width = 1.2; f.ws.column_dimensions["Z"].width = 4.2
    s = f.label(8, 26, 26, "نموذج ر ٣ – الضريبة على الرواتب والأجور – طبعة ٢٠١٠", 9, True, align=Alignment(horizontal="center", vertical="center", text_rotation=90), r2=36)
    s2 = f.label(40, 26, 26, "موقع وزارة المالية الالكتروني", 8, align=Alignment(horizontal="center", vertical="center", text_rotation=90), r2=LAST)
    return f.finish(path, LAST, 26)

def build_mof_r3_1(data=None, path=None):
    GREEN = "5DBF5D"; GL = "3A9A3A"
    f = Form("ر3-1", frame=GL, data=data); N = f.N
    f.fill(1, 4, 1, N, GREEN)
    f.label(1, 1, 7, "الجمهورية اللبنانيــــة", 10, True); f.label(2, 1, 7, "وزارة الماليـــــة", 10, True)
    f.label(3, 1, 7, "مديرية المالية العامة", 8); f.label(4, 1, 10, "مديرية الواردات – ضريبة الرواتب والأجور", 8)
    c = f.label(1, 8, 20, "كتاب طلب تســـجيل مســـتخدمين/أجـــراء", 18, True, align=center, r2=3); c.fill = PatternFill("solid", fgColor=GREEN)
    c = f.label(1, 21, N, "ر ١-٣", 18, True, align=center, r2=3); c.fill = PatternFill("solid", fgColor=GREEN)
    f.heights(range(1, 5), 15)
    r = 6; f.label(r, 1, 5, "إســـــم الشركة/المؤسسة", 10, True, True); f.field(r, 6, 13, key="company_name"); f.label(r, 14, 17, "الشهرة التجارية", 10); f.field(r, 18, N, key="trade_name")
    r = 8; f.label(r, 1, 6, "رقم تسجيل الشركة/المؤسسة", 10, True); f.boxes(r, 7, 10, key="company_mof")
    f.label(r, 17, 21, "رقــم الشـــركة/المؤسســة", 10, True); f.field(r, 22, N, key="company_nssf")
    f.label(r + 1, 1, 6, "(لدى وزارة المالية)", 9); f.label(r + 1, 17, 21, "(لدى الضمان الإجتماعي)", 9)
    f.hline(11)
    r = 12; f.label(r, 1, 6, "عنوان الشركة/المؤسسة", 10, True, True)
    r = 13
    for c1, t, c2 in ((1, "محافظة", 5), (7, "قضـــاء", 11), (13, "منطقة – بلدة", 17), (19, "الحـــي", N)):
        f.label(r, c1, c1 + 1, t, 9); f.field(r, c1 + 2, c2)
    r = 15; f.label(r, 1, 2, "الشـــارع", 9); f.field(r, 3, 8, key="company_address"); f.label(r, 9, 11, "المنطقة العقارية", 9); f.field(r, 12, 16)
    f.label(r, 17, 19, "رقم العقار / القسم", 9); f.field(r, 20, 23)
    r = 17; f.label(r, 1, 2, "المبنـــى", 9); f.field(r, 3, 7); f.label(r, 8, 9, "الطابـــق", 9); f.field(r, 10, 12)
    f.label(r, 13, 14, "هاتف", 9); f.field(r, 15, 18, key="company_phone"); f.label(r, 19, 20, "هاتف", 9); f.field(r, 21, N)
    r = 19; f.label(r, 1, 4, "صندوق بريد رقم", 9); f.field(r, 5, 8); f.label(r, 9, 10, "المنطقة", 9); f.field(r, 11, 16); f.label(r, 17, 18, "فاكس", 9); f.field(r, 19, 22)
    r = 21; f.label(r, 1, 6, "البريد الإلكتروني (e-mail):", 9); f.field(r, 7, 20, key="company_email")
    f.heights(range(6, 22), 20); f.hline(23)
    r = 24; f.label(r, 1, N, "في حال أرادت وزارة المالية الاستفسار عن أي موضوع يتعلق بطلبات تسجيل المستخدمين/الأجراء الرجاء الإتصال بـــ:", 10, True, True)
    r = 26; f.label(r, 1, 3, "الاسم الكامل", 9); f.field(r, 4, 13, key="contact_name"); f.label(r, 14, 15, "الصفة", 9); f.field(r, 16, 23, key="contact_title")
    r = 28; f.label(r, 1, 2, "هاتف", 9); f.field(r, 3, 7, key="company_phone"); f.label(r, 9, 10, "فاكس", 9); f.field(r, 11, 15)
    f.heights(range(24, 29), 20); f.hline(30)
    r = 32; f.label(r, 1, 13, "عدد طلبات تسجيل المستخدمين/الأجراء المرفقة بهذا الكتاب:", 10, True, True)
    cell = f.field(r, 15, 17, r + 1, box=True, key="requests"); cell.alignment = center
    f.hline(35)
    r = 36; f.label(r, 1, 4, "مقدم الكتاب", 10, True, True)
    r = 38; f.label(r, 1, 2, "الإسم", 9); f.field(r, 3, 12, key="contact_name"); f.label(r, 13, 14, "الصفة", 9); f.field(r, 15, 22, key="contact_title")
    r = 40; f.label(r, 1, 2, "التوقيع", 9); f.field(r, 3, 10); f.ws.row_dimensions[r].height = 30
    r = 42; f.label(r, 1, 2, "التاريخ", 9); f.date(r, 3, key="today")
    f.heights((32, 33, 36, 38, 42), 20)
    r = 45; f.fill(r, r, 1, N, "D9D9D9"); f.label(r, 1, N, "خاص بالإدارة", 10, True, align=center)
    r = 47; f.label(r, 1, 3, "رقم الكتاب", 9); f.boxes(r, 4, 10)
    f.label(r, 14, 16, "تـــاريخ تقديم الكتاب", 9); f.date(r, 17)
    LAST = 49; f.frame(1, LAST)
    f.ws.column_dimensions["Y"].width = 1.2; f.ws.column_dimensions["Z"].width = 4.2
    f.label(8, 26, 26, "نموذج ر ٣-١ – الضريبة على الرواتب والأجور – طبعة ٢٠١٠", 9, True, align=Alignment(horizontal="center", vertical="center", text_rotation=90), r2=30)
    f.label(36, 26, 26, "موقع وزارة المالية الالكتروني", 8, align=Alignment(horizontal="center", vertical="center", text_rotation=90), r2=LAST)
    return f.finish(path, LAST, 26)



FORMS = {
    "MOF_R3": ("ر٣ - طلب تسجيل مستخدم/أجير جديد (وزارة المالية)", build_mof_r3, True),
    "MOF_R3_1": ("ر٣-١ - كتاب طلب تسجيل مستخدمين/أجراء (وزارة المالية)", build_mof_r3_1, False),
    "CNSS_2AA": ("CNSS 2AA - تصريح باستخدام أجير", build_cnss_2aa, True),
    "CNSS_41A": ("CNSS 41A - إعلام عن استخدام أجير", build_cnss_41a, True),
    "CNSS_LEAVE": ("CNSS - إعلام عن ترك أجير عمله", build_cnss_leave, True),
}

MARITAL = {"single": "single", "married": "married", "spouse": "married", "widowed": "widowed", "widow": "widowed",
           "divorced": "divorced", "separated": "separated"}


PAY_TYPES = {"monthly": "monthly", "month": "monthly", "شهري": "monthly", "daily": "daily", "day": "daily", "يومي": "daily",
             "hourly": "hourly", "hour": "hourly", "بالساعة": "hourly", "weekly": "weekly", "اسبوعي": "weekly",
             "commission": "commission", "عمولة": "commission", "piece": "piece", "production": "piece", "على الانتاج": "piece"}
LEAVE_REASONS = ("استقالة | Resignation", "بلوغ السن | Retirement age", "عجز | Disability", "زواج | Marriage", "وفاة | Death",
                 "هجرة | Emigration", "عمل آخر | Other job")  # 2.9.108: the choices printed on the CNSS leaving notice, in order
LEAVE_WORDS = (("1", ("resign", "استقال")), ("2", ("retire", "السن", "تقاعد")), ("3", ("disab", "عجز")), ("4", ("marri", "زواج")),
               ("5", ("death", "died", "decease", "وفاة")), ("6", ("emigra", "travel", "هجرة", "سفر")), ("7", ("other", "job", "عمل")))


def leave_code(reason):
    """The printed code (1-7) of the leaving reason written in the employee file, or '' when it does not say."""
    text = str(reason or "").strip().lower()
    if text.isdigit() and text in "1234567": return text
    return next((code for code, words in LEAVE_WORDS if any(word in text for word in words)), "") if text else ""


def _display(value):
    text = str(value or "").strip()
    if len(text) >= 10 and text[4] == "-": return f"{text[8:10]}-{text[5:7]}-{text[:4]}"
    return text


NATIONALITIES_AR = {"lebanese": "لبنانية", "lebanon": "لبنانية", "syrian": "سورية", "syria": "سورية", "palestinian": "فلسطينية",
                    "egyptian": "مصرية", "iraqi": "عراقية", "jordanian": "أردنية", "sudanese": "سودانية", "ethiopian": "إثيوبية",
                    "filipino": "فلبينية", "philippine": "فلبينية", "bangladeshi": "بنغلادشية", "sri lankan": "سريلانكية", "srilankan": "سريلانكية",
                    "indian": "هندية", "nepalese": "نيبالية", "kenyan": "كينية", "ghanaian": "غانية", "nigerian": "نيجيرية", "french": "فرنسية",
                    "american": "أميركية", "british": "بريطانية", "canadian": "كندية", "turkish": "تركية", "armenian": "أرمنية", "saudi": "سعودية",
                    "emirati": "إماراتية", "kuwaiti": "كويتية", "qatari": "قطرية", "lebanese-american": "لبنانية"}


def arabic_nationality(value):
    """The nationality in Arabic: the Arabic text when it already is, a known English nationality translated, else as written."""
    text = str(value or "").strip()
    return NATIONALITIES_AR.get(text.lower(), text) if text else ""


def _ar(source, key, kind=None, language="ar"):
    """2.9.108: the Arabic version of a field (key + "_ar") when it is filled; otherwise (2.9.109) the English value written
    in Arabic when a kind is given (name, place, nationality, job), else the field as written.
    2.9.112: language "en" keeps the English value (the form is filled in English)."""
    if language == "en":
        return source.get(key) if str(source.get(key) or "").strip() else source.get(key + "_ar")
    value = str(source.get(key + "_ar") or "").strip()
    if value: return value
    if kind and str(source.get(key) or "").strip():
        from arabic_names import to_arabic
        return to_arabic(source.get(key), kind)
    return source.get(key)


def _money(value):
    try: return f"{float(str(value).replace(',', '')):,.0f}"
    except (TypeError, ValueError): return str(value or "")


EMPLOYEE_FORM_FIELDS = ("registry_number", "registry_place", "birth_caza", "sect", "work_hours", "work_time", "hire_salary", "work_permit",
                        "other_employer", "other_employer_nssf", "previous_employer", "previous_employer_nssf")  # 2.9.112: every box of the forms
WORK_TIMES = ("full", "part")


def form_data(company=None, employee=None, extra=None, language="ar"):
    """Values the forms use, from Settings > Company and the employee record. Nothing is invented:
    a value that is not in the program stays empty for the user to complete. Arabic fields are used first (2.9.108);
    language "en" (2.9.112) fills the forms with the English values instead."""
    company = company or {}; employee = employee or {}
    lang = "en" if str(language or "").lower().startswith("en") else "ar"
    pick = lambda source, key, kind=None: _ar(source, key, kind, lang)
    data = {key: company.get(key) for key in ("company_mof", "company_nssf", "company_phone", "company_email", "company_cr_number")}
    data.update(company_name=pick(company, "company_name", "company"), company_address=pick(company, "company_address", "place"),
                company_cr_place=pick(company, "company_cr_place", "place"), trade_name=company.get("company_trade_name") or "")
    today = date.today()
    data["today"] = today.strftime("%d-%m-%Y"); data["today_dm"] = today.strftime("%d-%m"); data["today_yy"] = today.strftime("%y")
    data["requests"] = "1"; data["language"] = lang; data["city"] = "بيروت" if lang == "ar" else "Beirut"
    if employee:
        name = str(pick(employee, "full_name", "name") or "").strip(); parts = name.split(None, 1)
        data.update({key: employee.get(key) for key in ("national_id", "nssf_number", "mof_number", "contact_number", "children", "registry_number",
                                                        "work_hours", "work_permit", "other_employer_nssf", "previous_employer_nssf")})
        data.update({key: pick(employee, key, kind) for key, kind in (("father_name", "name"), ("mother_name", "name"), ("birth_place", "place"),
                                                                      ("job_title", "job"), ("address", "place"), ("registry_place", "place"),
                                                                      ("birth_caza", "place"), ("sect", None), ("other_employer", "name"),
                                                                      ("previous_employer", "name"))})
        # 2.9.109: spouse (MOF R3, CNSS 2AA family page)
        data.update({key: pick(employee, key, kind) for key, kind in (("spouse_name", "name"), ("spouse_maiden_name", "name"), ("spouse_father", "name"),
                                                                      ("spouse_mother", "name"), ("spouse_nationality", "nationality"),
                                                                      ("spouse_birth_place", "place"))})
        data.update({key: employee.get(key) for key in ("spouse_id", "spouse_mof", "spouse_employer", "spouse_employer_mof", "spouse_public_admin",
                                                        "family_dependents")})
        data["spouse_birth_date"] = _display(employee.get("spouse_birth_date"))
        data["spouse_birth"] = " - ".join(filter(None, [data.get("spouse_birth_place"), data["spouse_birth_date"]]))
        data["spouse_full"] = " ".join(filter(None, [str(data.get("spouse_name") or "").strip(), str(data.get("spouse_maiden_name") or "").strip()]))
        data.update(full_name=name, first_name=parts[0] if parts else "", last_name=parts[1] if len(parts) > 1 else "",
                    birth_date=_display(employee.get("birth_date")), hire_date=_display(employee.get("hire_date")),
                    leave_date=_display(employee.get("leave_date")),
                    birth=" - ".join(filter(None, [data.get("birth_place"), _display(employee.get("birth_date"))])),
                    nationality=pick(employee, "nationality", "nationality") or "",
                    marital_status=MARITAL.get(str(employee.get("marital_status") or "").lower(), ""),
                    spouse_works=("yes" if int(employee.get("spouse_works") or 0) else "no") if str(employee.get("marital_status") or "").lower() in ("married", "spouse") else "",
                    has_mof="yes" if employee.get("mof_number") else "no", sex=str(employee.get("sex") or "").lower(),
                    pay_type=PAY_TYPES.get(str(employee.get("pay_type") or "monthly").strip().lower(), "monthly"),
                    leave_code=leave_code(employee.get("leave_reason")),
                    work_time=str(employee.get("work_time") or "").strip().lower() if str(employee.get("work_time") or "").strip().lower() in WORK_TIMES else "",
                    other_job="yes" if str(employee.get("other_employer") or "").strip() else "no")
        # 2.9.112: the address of the employee, box by box (governorate falls back on the one-line address)
        for key in ("addr_governorate", "addr_caza", "addr_town", "addr_district", "addr_street", "addr_building"):
            if str(employee.get(key) or employee.get(key + "_ar") or "").strip(): data[key] = pick(employee, key, "place")
        data.update({key: employee.get(key) for key in ("addr_floor", "phone2") if employee.get(key)})
        if not data.get("addr_governorate") and data.get("address"): data["addr_governorate"] = data["address"]
        data["building_floor"] = " - ".join(filter(None, [str(data.get("addr_building") or "").strip(),
                                                          (("طابق " if lang == "ar" else "Floor ") + str(data["addr_floor"])) if data.get("addr_floor") else ""]))
        nationality = str(data.get("nationality") or "").strip().lower()
        if nationality and not any(t in nationality for t in ("leban", "لبنان")): data["foreign_nationality"] = data["nationality"]
        lbp = str(employee.get("currency") or "LBP").upper() == "LBP"
        if lbp and employee.get("base_salary"):
            data["salary_lbp"] = _money(employee["base_salary"])
        hire_salary = employee.get("hire_salary") or employee.get("base_salary")  # 2.9.112: wage at the date of entry (the salary when not given)
        if lbp and hire_salary: data["hire_salary"] = _money(hire_salary)
        if employee.get("base_salary"):
            currency = {"LBP": "ل.ل.", "USD": "د.أ."}.get(str(employee.get("currency") or "LBP").upper(), employee.get("currency") or "")
            data["salary_text"] = f"{float(employee['base_salary']):,.0f} {currency}"
    data.update(extra or {})
    return data


EDITABLE_FIELDS = (  # 2.9.112: what the user can correct on the screen before the Excel is made (label, key)
    ("Company name | إسم المؤسسة", "company_name"), ("Company address | عنوان المؤسسة", "company_address"),
    ("Commercial register no. | رقم السجل التجاري", "company_cr_number"), ("Commercial register place | مكان السجل التجاري", "company_cr_place"),
    ("Company phone | هاتف", "company_phone"), ("Employer NSSF no. | رقم المؤسسة في الضمان", "company_nssf"),
    ("Company MOF no. | رقم التسجيل في المالية", "company_mof"), ("Trade name | الشهرة التجارية", "trade_name"),
    ("First name | الإسم", "first_name"), ("Family name | الشهرة", "last_name"), ("Full name | الإسم الكامل", "full_name"), ("Father | إسم الأب", "father_name"),
    ("Mother | إسم الأم وشهرتها", "mother_name"), ("Place of birth | محل الولادة", "birth_place"), ("Date of birth | تاريخ الولادة", "birth_date"),
    ("Birth - place and date | تاريخ ومحل الولادة", "birth"), ("Caza of birth | القضاء", "birth_caza"),
    ("Registry no. | رقم السجل", "registry_number"), ("Registry place | مكان السجل / محل الإقامة حسب الهوية", "registry_place"),
    ("Sect | المذهب", "sect"), ("Nationality | الجنسية", "nationality"), ("ID card no. | رقم بطاقة الهوية", "national_id"),
    ("Employee NSSF no. | رقم الأجير في الضمان", "nssf_number"), ("Employee MOF no. | الرقم المالي", "mof_number"),
    ("Governorate | المحافظة", "addr_governorate"), ("Caza | القضاء", "addr_caza"), ("Town | المدينة أو القرية", "addr_town"),
    ("District | الحي", "addr_district"), ("Street | الشارع", "addr_street"), ("Building | بناية", "addr_building"), ("Floor | طابق", "addr_floor"),
    ("Building and floor | بناية وطابق", "building_floor"), ("Phone | هاتف", "contact_number"),
    ("Hire date | تاريخ دخول العمل", "hire_date"), ("Hours per month | عدد ساعات العمل في الشهر", "work_hours"),
    ("Job | عمل الأجير الحالي", "job_title"), ("Current salary | الراتب الحالي", "salary_lbp"), ("Salary at hire | الاجر بتاريخ دخول العمل", "hire_salary"),
    ("Salary text | الراتب", "salary_text"), ("Work permit no. and date | رقم وتاريخ اجازة العمل", "work_permit"),
    ("Other employer | صاحب العمل الآخر", "other_employer"), ("Other employer NSSF no. | رقمه", "other_employer_nssf"),
    ("Previous employer | صاحب العمل السابق", "previous_employer"), ("Previous employer NSSF no. | رقمه", "previous_employer_nssf"),
    ("Leave date | تاريخ ترك العمل", "leave_date"), ("Spouse | الزوج/الزوجة", "spouse_full"), ("Children | عدد الأولاد", "children"),
    ("Date on the form | التاريخ", "today"))


def build_form(key, path, company=None, employee=None, extra=None, language="ar"):
    title, builder, needs_employee = FORMS[key]
    if needs_employee and not employee: raise ValueError("Select an employee first")
    return builder(form_data(company, employee, extra, language), path)
