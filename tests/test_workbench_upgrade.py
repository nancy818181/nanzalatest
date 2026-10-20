"""Daily-workflow regression tests: real SQLite books, scoped HTTP and UI concurrency."""

import tempfile
import threading
import time
import unittest
from datetime import datetime
from decimal import Decimal
from http.server import ThreadingHTTPServer
from pathlib import Path
from unittest.mock import patch, Mock

from database import Database
from client import ApiClient
from server import ApiHandler
from company_manager import CompanyManager
from nanza_delivery.ageing_report import bucket_for, get_ageing_buckets
from nanza_delivery.whatsapp_daily_summary import (
    generate_whatsapp_summary,
    reminder_due,
    preferences,
    save_preferences,
)
from nanza_delivery.custom_report_builder import (
    build_report,
    export_report,
    save_template,
    list_templates,
    validate_filters,
)
from nanza_delivery.ai_accountant import (
    suggest_category,
    detect_duplicate,
    cashflow_forecast,
    explain_profit_loss,
)


class BucketTest(unittest.TestCase):
    def test_exact_boundaries(self):
        for days, expected in (
            (0, "0-30"),
            (30, "0-30"),
            (31, "31-60"),
            (60, "31-60"),
            (61, "61-90"),
            (90, "61-90"),
            (91, "90+"),
        ):
            with self.subTest(days=days):
                self.assertEqual(bucket_for(days), expected)

    def test_empty_requires_company(self):
        with self.assertRaises(ValueError):
            get_ageing_buckets("2026-10-01")

    def test_reminder_never_sends_and_fires_once(self):
        self.assertFalse(reminder_due(datetime(2026, 10, 18, 17, 59), True))
        self.assertTrue(reminder_due(datetime(2026, 10, 18, 18, 0), True))
        self.assertFalse(reminder_due(datetime(2026, 10, 18, 19, 0), True, "2026-10-18"))
        self.assertFalse(reminder_due(datetime(2026, 10, 18, 19, 0), False))

    def test_classifier_abstains_without_active_company_accounts(self):
        self.assertTrue(suggest_category("fuel")["abstain"])
        accounts = [
            {"code": "699900001", "name_en": "Fuel", "type": "expense", "is_active": 0, "is_posting": 1}
        ]
        self.assertIsNone(suggest_category("fuel", accounts)["account_code"])

    def test_classifier_uses_actual_chart_and_ambiguity(self):
        account = {"code": "626340001", "name_en": "Fuel", "type": "expense", "is_active": 1, "is_posting": 1}
        self.assertEqual(suggest_category("fuel", [account])["account_code"], "626340001")
        self.assertIsNone(suggest_category("fuel", [account, dict(account, code="626340002")])["account_code"])
        self.assertIsNone(suggest_category("fuel", [dict(account, code="X-EXP-19")])["account_code"])

    def test_sql_and_invalid_ranges_refused(self):
        for filters in (
            {"sql": "DROP TABLE invoices"},
            {"date_from": "2026-02-30", "date_to": "2026-03-01"},
            {"date_from": "2026-02-01", "date_to": "2026-01-01"},
            {"date_from": "2025-12-01", "date_to": "2026-01-01"},
            {"date_from": "2026-01-01", "date_to": "2026-02-01", "customer": "1 OR 1=1"},
        ):
            with self.subTest(filters=filters), self.assertRaises(ValueError):
                validate_filters(filters)


class WorkbenchDatabaseTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.db = Database(Path(self.tmp.name) / "company.db")
        self.db.initialize("Workbench-test-2026!")
        self.addCleanup(self.db.release)
        self.user = 1
        self.filters = {"date_from": "2026-01-01", "date_to": "2026-12-31"}

    def invoice(
        self,
        number="SALE-1",
        amount=100,
        currency="USD",
        day="2026-01-01",
        status="posted",
        subtype="invoice",
        party="Client A",
        line=True,
    ):
        if subtype == "credit_note":
            return self.db.create_manual_invoice(
                {
                    "invoice_number": number,
                    "invoice_date": day,
                    "due_date": day,
                    "kind": "sales",
                    "party_name": party,
                    "currency": currency,
                    "status": status,
                    "doc_subtype": subtype,
                    "supplier_side": "C",
                    "vat_side": "D",
                    "expense_side": "D",
                },
                [
                    {
                        "description": "منتج Product",
                        "quantity": 2,
                        "unit_price": amount / 2,
                        "vat_rate": 0,
                        "item_code": "SKU-A",
                    }
                ],
                1,
            )
        saved = self.db.import_invoice(
            {
                "invoice_number": number,
                "invoice_date": day,
                "due_date": day,
                "kind": "sale",
                "party_name": party,
                "currency": currency,
                "subtotal": amount,
                "vat": 0,
                "total": amount,
                "status": status,
                "doc_subtype": subtype,
            },
            self.user,
        )
        if line:
            with self.db.connect() as c:
                c.execute(
                    """INSERT INTO invoice_items(invoice_id,description,quantity,unit_price,subtotal,vat_rate,vat,total,item_code)
                    VALUES(?, 'منتج Product', '2', ?, ?, '0', '0', ?, 'SKU-A')""",
                    (saved, str(amount / 2), str(amount), str(amount)),
                )
        return saved

    def voucher(self, day, debit, credit, amount):
        return self.db.save_journal_voucher(
            {"entry_date": day, "currency": "USD", "description": "Workbench"},
            [{"account_code": debit, "debit": amount}, {"account_code": credit, "credit": amount}],
            self.user,
        )

    def test_ageing_native_currencies_and_credits_reconcile(self):
        self.invoice(amount=100)
        self.invoice("LBP-1", 10000, "LBP")
        report = get_ageing_buckets("2026-02-01", self.db)
        self.assertEqual(report["totals"]["USD"]["total"], Decimal("100"))
        self.assertEqual(report["totals"]["LBP"]["total"], Decimal("10000"))
        self.assertEqual(report["totals"]["USD"]["31-60"], Decimal("100"))
        self.assertEqual(sum(r["total"] for r in report["parties"] if r["currency"] == "USD"), Decimal("100"))

    def test_ageing_historical_payments_and_customer_drilldown(self):
        invoice = self.invoice()
        party = next(p for p in self.db.list_parties() if p["name"] == "Client A")
        payment = self.db.add_payment(
            {
                "kind": "customer_receipt",
                "payment_date": "2026-02-01",
                "party_id": party["id"],
                "currency": "USD",
                "amount": 40,
                "cash_account": "531",
                "party_account": party["account_number"],
            },
            1,
        )
        self.db.save_allocations(payment, [{"invoice_id": invoice, "amount": 40}], 1)
        before = get_ageing_buckets("2026-01-31", self.db)
        after = get_ageing_buckets("2026-02-02", self.db)
        self.assertEqual(before["totals"]["USD"]["total"], Decimal("100"))
        self.assertEqual(after["parties"][0]["invoices"][0]["outstanding"], Decimal("60"))

    def test_custom_report_exact_amounts_and_signed_notes(self):
        self.invoice(amount=111)
        self.invoice("CN-1", 11, subtype="credit_note")
        report = build_report(self.db, self.filters)
        self.assertEqual(report["totals"]["USD"], Decimal("100"))
        self.assertEqual(len(report["rows"]), 2)
        self.assertTrue(any(r[-1] == Decimal("-11") for r in report["rows"]))

    def test_custom_report_filters_and_no_header_inflation(self):
        self.invoice()
        self.invoice("OTHER", 50, party="Client B")
        self.invoice("HEADER", 999, line=False)
        party = next(p for p in self.db.list_parties() if p["name"] == "Client A")
        report = build_report(self.db, dict(self.filters, customer=party["id"], product="SKU-A"))
        self.assertEqual(len(report["rows"]), 1)
        self.assertEqual(report["totals"]["USD"], Decimal("100"))

    def test_review_invoices_not_reported(self):
        self.invoice(status="review")
        self.assertEqual(build_report(self.db, self.filters)["rows"], [])
        self.assertNotIn("100.00", generate_whatsapp_summary("2026-01-01", self.db))

    def test_warehouse_requires_real_link_and_does_not_multiply(self):
        saved = self.invoice()
        with self.db.connect() as c:
            line = c.execute("SELECT id FROM invoice_items WHERE invoice_id=?", (saved,)).fetchone()[0]
            item = c.execute("INSERT INTO inventory_items(sku,name) VALUES('SKU-A','Test')").lastrowid
            for _ in range(2):
                c.execute(
                    """INSERT INTO stock_movements(item_id,movement_date,quantity,unit_cost,warehouse_id,invoice_item_id)
                    VALUES(?,'2026-01-01','1','2',1,?)""",
                    (item, line),
                )
        report = build_report(self.db, dict(self.filters, warehouse="1"))
        self.assertEqual(report["totals"]["USD"], Decimal("100"))
        self.assertEqual(build_report(self.db, dict(self.filters, warehouse="999"))["rows"], [])

    def test_pagination_and_complete_export(self):
        for i in range(3):
            self.invoice(str(i), i + 1)
        first = build_report(self.db, self.filters, limit=2)
        self.assertTrue(first["has_more"])
        self.assertEqual(len(build_report(self.db, self.filters, limit=2, offset=2)["rows"]), 1)
        full = export_report(self.db, self.filters)
        self.assertEqual(len(full["rows"]), 3)
        self.assertEqual(full["totals"]["USD"], Decimal("6"))
        with self.assertRaisesRegex(ValueError, "narrow"):
            export_report(self.db, self.filters, max_rows=2)

    def test_templates_get_is_read_only_and_user_isolation(self):
        self.assertEqual(list_templates(self.db, 1), [])
        with self.db.connect() as c:
            self.assertIsNone(c.execute("SELECT 1 FROM sqlite_master WHERE name='custom_reports'").fetchone())
        save_template(self.db, 1, "My Monthly Sales", self.filters)
        self.assertEqual(list_templates(self.db, 2), [])
        self.assertEqual(list_templates(self.db, 1)[0]["filters"], self.filters)
        save_template(self.db, 1, "My Monthly Sales", dict(self.filters, currency="LBP"))
        self.assertEqual(list_templates(self.db, 1)[0]["revision"], 2)

    def test_summary_signed_sales_cash_and_bilingual(self):
        self.invoice(amount=100)
        self.invoice("CN-1", 10, subtype="credit_note")
        self.invoice("OTHER-CURRENCY", 10000, "LBP")
        self.voucher("2026-01-01", "531", "10", 500)
        text = generate_whatsapp_summary("2026-01-01", self.db)
        self.assertIn("Posted sales today (VAT included): 90.00", text)
        self.assertIn("balance: 500.00", text)
        self.assertNotIn("10,000.00", text)
        self.assertIn("ملخص نانزا", generate_whatsapp_summary("2026-01-01", self.db, language="ar"))

    def test_duplicate_before_save_and_edit_exclusion(self):
        invoice = self.invoice("DUP-1", 111)
        self.assertEqual(len(detect_duplicate("DUP-1", 111, "2026-01-01", self.db, currency="USD")), 1)
        self.assertEqual(detect_duplicate("DUP-1", 111, "2026-01-01", self.db, exclude_id=invoice), [])
        self.assertEqual(detect_duplicate("DUP-1", 111, "2026-01-01", self.db, currency="LBP"), [])
        self.assertEqual(detect_duplicate("DUP-1", 112, "2026-01-01", self.db), [])
        self.assertEqual(detect_duplicate("", 111, "2026-01-01", self.db), [])

    def test_forecast_no_history_abstains(self):
        report = cashflow_forecast(self.db, as_of="2026-01-31")
        self.assertFalse(report["available"])
        self.assertEqual(report["points"], [])

    def test_forecast_uses_cash_not_accrual_and_no_opening_growth(self):
        self.voucher("2026-01-01", "531", "10", 1000)
        self.voucher("2026-01-15", "531", "7011", 310)
        self.invoice("UNPAID", 10000, day="2026-01-20")
        report = cashflow_forecast(self.db, as_of="2026-01-31", history_days=31)
        # Journal vouchers are actual cash movements; even a capital injection is a historical cash flow.
        self.assertEqual(report["opening_cash"], Decimal("1310"))
        self.assertEqual(report["daily_average"], Decimal("1310") / 31)
        self.assertEqual(len(report["points"]), 31)
        self.assertNotIn("10000", str(report["points"]))

    def test_profit_explanation_computes_actual_comparison(self):
        self.voucher("2026-02-01", "531", "7011", 200)
        self.voucher("2026-02-02", "6263.1", "531", 20)
        self.voucher("2026-03-01", "531", "7011", 200)
        self.voucher("2026-03-02", "6263.1", "531", 50)
        # March 1-28 versus February 1-28.
        text = explain_profit_loss(self.db, start="2026-03-01", end="2026-03-28")
        self.assertIn("change -30.00", text)
        self.assertIn("30.00", text)
        self.assertIn("not proof of causation", text)
        self.assertIn(
            "صافي الربح", explain_profit_loss(self.db, start="2026-03-01", end="2026-03-28", language="ar")
        )

    def test_no_fabricated_prior_year_explanation(self):
        self.assertIn("prior-year", explain_profit_loss(self.db, start="2026-01-01", end="2026-01-31"))

    def test_reminder_preferences_scoped_and_no_transmission(self):
        self.assertFalse(preferences(self.db, 1)["reminder_6pm"])
        save_preferences(self.db, 1, {"reminder_6pm": True, "last_day": "2026-10-18"})
        self.assertTrue(preferences(self.db, 1)["reminder_6pm"])
        self.assertFalse(preferences(self.db, 2)["reminder_6pm"])
        self.assertFalse(preferences(self.db, 1)["auto_send_available"])


class WorkbenchHttpTest(unittest.TestCase):
    def test_auth_company_and_viewer_guards(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as folder:
            path = Path(folder) / "master.db"
            db = Database(path)
            db.initialize("Http-workbench-test!")
            manager = CompanyManager(path)
            if not manager.list_companies():
                manager.create_company({"name": "Workbench Sample", "year": 2026}, db)
            handler = type(
                "WorkbenchHandler", (ApiHandler,), dict(db=db, master_db=db, company_manager=manager)
            )
            http = ThreadingHTTPServer(("127.0.0.1", 0), handler)
            thread = threading.Thread(target=http.serve_forever, daemon=True)
            thread.start()
            try:
                api = ApiClient(f"http://127.0.0.1:{http.server_port}")
                with self.assertRaises(RuntimeError):
                    api.workbench("templates")
                api.login("admin", "Http-workbench-test!")
                company = api.companies()[0]
                year = int(company["years"][0]["year"])
                api.select_company_year(company["id"], year)
                filters = {"date_from": f"{year}-01-01", "date_to": f"{year}-12-31"}
                api.save_custom_report("Scoped report", filters)
                self.assertEqual(len(api.workbench("templates")["items"]), 1)
                self.assertEqual(api.workbench("report", **filters)["rows"], [])
                self.assertIn("parties", api.workbench("ageing", date=f"{year}-12-31"))
                self.assertIn("text", api.workbench("summary", date=f"{year}-12-31"))
                self.assertFalse(api.workbench("forecast", date=f"{year}-12-31")["available"])
                self.assertTrue(api.workbench("suggest", description="unknown")["abstain"])
                self.assertEqual(
                    api.workbench("duplicates", invoice_no="X", amount="1", date=f"{year}-01-01")["items"], []
                )
                db.save_user({"username": "wbviewer", "password": "Viewer-workbench!", "role": "viewer"}, 1)
                viewer = ApiClient(api.base_url)
                viewer.login("wbviewer", "Viewer-workbench!")
                viewer.select_company_year(company["id"], year)
                self.assertEqual(viewer.workbench("templates")["items"], [])
                with self.assertRaisesRegex(RuntimeError, "read-only"):
                    viewer.save_custom_report("Forbidden", filters)
                manager.create_year(company["id"], year + 1, 1)
                api.select_company_year(company["id"], year + 1)
                self.assertEqual(api.workbench("templates")["items"], [])
                api.close()
                viewer.close()
            finally:
                http.shutdown()
                http.server_close()
                thread.join(5)
                for cached in list(manager._cache):
                    manager._cache.pop(cached).release()
                db.release()


if __name__ == "__main__":
    unittest.main()


class WorkbenchGuiTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import tkinter as tk

        try:
            probe = tk.Tk()
            probe.destroy()
        except tk.TclError:
            raise unittest.SkipTest("No graphical display available")

    def setUp(self):
        import tkinter as tk
        from desktop_workbench import WorkbenchMixin

        self.tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.addCleanup(self.tmp.cleanup)
        path = Path(self.tmp.name) / "master.db"
        self.db = Database(path)
        self.db.initialize("Gui-workbench-test!")
        self.manager = CompanyManager(path)
        if not self.manager.list_companies():
            self.manager.create_company({"name": "GUI company", "year": 2026}, self.db)
        company = self.manager.list_companies()[0]
        self.year = int(company["years"][0]["year"])
        self.books = self.manager.database(company["id"], self.year)
        self.books.import_invoice(
            {
                "invoice_number": "GUI-SALE",
                "invoice_date": f"{self.year}-01-01",
                "due_date": f"{self.year}-01-01",
                "kind": "sale",
                "party_name": "GUI Customer",
                "currency": "USD",
                "subtotal": 100,
                "vat": 0,
                "total": 100,
            },
            1,
        )
        with self.books.connect() as c:
            invoice = c.execute("SELECT id FROM invoices WHERE invoice_number='GUI-SALE'").fetchone()[0]
            c.execute(
                """INSERT INTO invoice_items(invoice_id,description,quantity,unit_price,subtotal,vat_rate,vat,total,item_code)
                VALUES(?,'اختبار Test','1','100','100','0','0','100','GUI-SKU')""",
                (invoice,),
            )
        handler = type(
            "GuiWorkbenchHandler",
            (ApiHandler,),
            dict(db=self.db, master_db=self.db, company_manager=self.manager),
        )
        self.http = ThreadingHTTPServer(("127.0.0.1", 0), handler)
        self.thread = threading.Thread(target=self.http.serve_forever, daemon=True)
        self.thread.start()

        class TestApp(WorkbenchMixin, tk.Tk):
            def fiscal_today(app):
                return f"31-01-{self.year}"

        self.app = TestApp()
        self.app.geometry("1366x768")
        self.app.client = ApiClient(f"http://127.0.0.1:{self.http.server_port}")
        self.app.client.login("admin", "Gui-workbench-test!")
        self.app.client.select_company_year(company["id"], self.year)
        self.app.language = tk.StringVar(value="en")
        self.app.dashboard_base_currency = tk.StringVar(value="USD")
        import arabic_ui

        arabic_ui.switch(self.app, "en")
        self.app.errors = []
        self.app.report_callback_exception = lambda *args: self.app.errors.append(args)
        self.dialog_patch = patch(
            "desktop_workbench.messagebox.showerror", side_effect=lambda *a, **k: self.app.errors.append(a)
        )
        self.dialog_patch.start()
        self.addCleanup(self.dialog_patch.stop)

    def tearDown(self):
        import arabic_ui

        arabic_ui.switch(self.app, "en")
        self.app.destroy()
        self.app.client.close()
        self.http.shutdown()
        self.http.server_close()
        self.thread.join(5)
        for cached in list(self.manager._cache):
            self.manager._cache.pop(cached).release()
        self.db.release()

    def wait(self, predicate, timeout=8):
        deadline = time.monotonic() + timeout
        while time.monotonic() < deadline:
            self.app.update()
            if predicate():
                return
            time.sleep(0.01)
        self.fail("UI did not complete its background operation")

    def test_ageing_filter_drilldown_and_excel_export(self):
        from desktop_workbench import AgeingWindow

        window = AgeingWindow(self.app)
        self.wait(lambda: window.data is not None)
        self.assertEqual(len(window.tree.get_children()), 1)
        window.filter("90+")
        self.assertEqual(len(window.tree.get_children()), 0)
        window.filter("0-30")
        window.tree.selection_set("0")
        window.open_invoices()
        self.assertGreaterEqual(len(window.winfo_children()), 4)
        output = Path(self.tmp.name) / "ageing.xlsx"
        with patch("desktop_workbench.filedialog.asksaveasfilename", return_value=str(output)):
            window.export()
            self.wait(lambda: getattr(window, "_wb_export_done", False))
        from openpyxl import load_workbook

        book = load_workbook(output)
        self.assertIn("GUI Customer", [cell.value for row in book.active for cell in row])
        book.close()
        self.assertEqual(self.app.errors, [])

    def test_custom_live_preview_template_pdf_excel_and_rtl(self):
        from desktop_workbench import CustomReportWindow

        self.app.language.set("ar")
        import arabic_ui

        arabic_ui.switch(self.app, "ar")
        window = CustomReportWindow(self.app)
        self.wait(lambda: window.result is not None and bool(window.matches))
        self.assertEqual(len(window.tree.get_children()), 1)
        window.vars["currency"].set("LBP")
        self.wait(lambda: window.result is not None and window.result["filters"].get("currency") == "LBP")
        self.assertEqual(window.result["rows"], [])
        window.vars["currency"].set("USD")
        self.wait(lambda: window.result is not None and window.result["filters"].get("currency") == "USD")
        with patch("desktop_workbench.simpledialog.askstring", return_value="My Monthly Sales"):
            window.save_template()
        self.wait(lambda: "My Monthly Sales" in window.templates)
        for extension in ("xlsx", "pdf"):
            output = Path(self.tmp.name) / ("custom." + extension)
            with patch("desktop_workbench.filedialog.asksaveasfilename", return_value=str(output)):
                window.export(extension)
                self.wait(lambda: output.exists() and getattr(window, "_wb_export_done", False))
            self.assertGreater(output.stat().st_size, 500)
            if extension == "xlsx":
                from openpyxl import load_workbook

                book = load_workbook(output)
                self.assertTrue(book.active.sheet_view.rightToLeft)
                book.close()
        self.assertEqual(self.app.errors, [])

    def test_whatsapp_preview_copy_and_dashboard_card(self):
        from desktop_workbench import WhatsAppWindow
        import tkinter as tk

        panel = tk.Frame(self.app)
        panel.pack(fill="x")
        self.app.build_workbench_tools(panel)
        window = WhatsAppWindow(self.app)
        self.wait(lambda: window.ready)
        window.copy()
        self.assertIn("Posted sales today", window.clipboard_get())
        self.assertEqual(str(window.copy_button["state"]), "normal")
        window.currency.set("LBP")
        self.assertFalse(window.ready)
        self.assertEqual(str(window.copy_button["state"]), "disabled")
        self.wait(lambda: "USD" in self.app._wb_badge["text"])
        self.assertEqual(self.app.errors, [])

    def test_stale_results_discarded_on_company_switch(self):
        import tkinter as tk
        from desktop_workbench import async_call

        owner = tk.Frame(self.app)
        owner.pack()
        release = threading.Event()
        received = []
        threads = []

        def work(_client):
            threads.append(threading.get_ident())
            release.wait(2)
            return "old company"

        async_call(owner, self.app, work, received.append)
        self.app.client.company_id = "different-company"
        release.set()
        self.wait(lambda: bool(threads))
        deadline = time.monotonic() + 0.2
        while time.monotonic() < deadline:
            self.app.update()
            time.sleep(0.01)
        self.assertEqual(received, [])
        self.assertNotEqual(threads[0], threading.get_ident())

    def test_closing_window_cancels_scheduled_callbacks(self):
        from desktop_workbench import CustomReportWindow

        window = CustomReportWindow(self.app)
        self.wait(lambda: window.result is not None and bool(window.matches))
        timer = window._wb_poll_id
        scope_timer = window._wb_scope_id
        window.destroy()
        remaining = self.app.tk.call("after", "info")
        self.assertNotIn(timer, remaining)
        self.assertNotIn(scope_timer, remaining)
        self.app.update()
        self.assertEqual(self.app.errors, [])
