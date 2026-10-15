"""2.9.119: exports (Excel / PDF / printed invoice) in Arabic, right to left, when the program is in Arabic; English unchanged."""
import tests_setup  # noqa: F401
import tempfile
import unittest
from pathlib import Path

import arabic_ui
import report_export

SECTIONS = [{"heading": "Trial Balance", "headers": ["Account", "Name", "Debit", "Credit"], "rows": [["5121", "Bank", 10.0, 0.0], ["", "TOTAL", 10.0, 0.0]],
             "total_rows": [1]}]
INVOICE = {"invoice_number": "SAL-1", "invoice_date": "15-10-2026", "currency": "USD", "kind": "sale", "party_name": "Cedar", "subtotal": 100,
           "vat": 11, "total": 111, "amount_paid": 0}


class ArabicReportsTest(unittest.TestCase):
    def setUp(self):
        self.folder = Path(tempfile.mkdtemp()); arabic_ui.install()
        self.addCleanup(lambda: setattr(arabic_ui, "ON", False))

    def excel(self, name):
        from openpyxl import load_workbook
        path = self.folder / name; report_export.export_sections_excel(path, "Trial Balance", ["Period 01-01-2026 - 31-12-2026"], SECTIONS)
        return load_workbook(path).active

    def test_english_unchanged(self):
        arabic_ui.ON = False
        sheet = self.excel("en.xlsx")
        self.assertFalse(sheet.sheet_view.rightToLeft)
        self.assertEqual(sheet.cell(1, 1).value, "Trial Balance")

    def test_arabic_excel_and_pdf(self):
        arabic_ui.ON = True
        sheet = self.excel("ar.xlsx")
        self.assertTrue(sheet.sheet_view.rightToLeft)
        values = [c.value for row in sheet.iter_rows() for c in row if c.value]
        self.assertIn("ميزان المراجعة", values)
        self.assertIn(arabic_ui.translate("Debit"), values)
        self.assertIn("Bank", values)  # names and amounts stay as they are
        pdf = self.folder / "ar.pdf"; report_export.export_sections_pdf(pdf, "Trial Balance", [], SECTIONS)
        invoice = self.folder / "invoice.pdf"
        report_export.export_invoice_pdf(invoice, INVOICE, [{"item_code": "A", "description": "Service", "quantity": 1, "unit_price": 100}], None, {"company_name": "Cedar"})
        for path in (pdf, invoice):
            data = path.read_bytes(); self.assertTrue(data.startswith(b"%PDF")); self.assertIn(b"Amiri", data)


if __name__ == "__main__":
    unittest.main()
