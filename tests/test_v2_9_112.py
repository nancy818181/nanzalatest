"""2.9.112: every box of the NSSF / MOF forms filled, the forms in Arabic or English, and the P&L on the tax basis."""
import tempfile
import unittest
from pathlib import Path

import payroll_excel_forms as forms
from database import Database
from openpyxl import load_workbook

COMPANY = {"company_name": "Saber for Audit", "company_name_ar": "", "company_nssf": "601123", "company_cr_number": "12345",
           "company_cr_place": "Baabda", "company_address": "Hamra Street", "company_phone": "01-123456"}
EMPLOYEE = {"full_name": "Bassam Saber", "father_name": "Fouad", "mother_name": "Mary", "birth_date": "1986-01-19", "birth_place": "Beirut",
            "nationality": "Lebanese", "job_title": "Sales manager", "base_salary": "35000000", "currency": "LBP", "hire_date": "2025-11-01",
            "registry_number": "45", "registry_place": "Zouk Mikael", "birth_caza": "Kesrouan", "work_hours": "180", "work_time": "full",
            "hire_salary": "30000000", "addr_governorate": "Mount Lebanon", "addr_caza": "Kesrouan", "addr_town": "Zouk Mikael",
            "addr_street": "Main Road", "addr_building": "Sayegh", "addr_floor": "3", "contact_number": "03-111222", "sex": "male",
            "marital_status": "married", "leave_reason": "Resignation", "leave_date": "2026-10-10"}


def texts(path):
    return [str(c.value) for ws in load_workbook(path).worksheets for row in ws.iter_rows() for c in row if c.value not in (None, "")]


class FormsEveryBoxTest(unittest.TestCase):
    def test_new_boxes_are_filled_in_arabic(self):
        data = forms.form_data(COMPANY, EMPLOYEE)
        self.assertEqual(data["company_name"], "صابر للتدقيق")
        self.assertEqual(data["company_address"], "شارع الحمرا")
        self.assertEqual(data["hire_salary"], "30,000,000")
        self.assertEqual(data["registry_number"], "45")
        self.assertEqual(data["registry_place"], "ذوق مكايل")
        self.assertEqual(data["work_time"], "full")
        self.assertEqual(data["other_job"], "no")
        self.assertIn("طابق 3", data["building_floor"])
        with tempfile.TemporaryDirectory() as folder:
            for key in ("CNSS_2AA", "CNSS_LEAVE", "CNSS_41A", "MOF_R3"):
                with self.subTest(form=key):
                    path = Path(folder) / f"{key}.xlsx"; forms.build_form(key, path, COMPANY, EMPLOYEE); text = texts(path)
                    self.assertIn("صابر للتدقيق", text)
                    if key != "MOF_R3": self.assertIn("12345", text)       # commercial register on the NSSF forms
                    self.assertIn("45", text)                              # registry number
                    if key == "CNSS_2AA":
                        for value in ("30,000,000", "ذوق مكايل", "كسروان", "180", data["today"]): self.assertIn(value, text)
                    if key == "CNSS_LEAVE": self.assertIn("مدير مبيعات", text)  # current job, empty before

    def test_english_keeps_the_english_values(self):
        data = forms.form_data(COMPANY, EMPLOYEE, language="en")
        self.assertEqual(data["company_name"], "Saber for Audit")
        self.assertEqual(data["first_name"], "Bassam")
        self.assertEqual(data["nationality"], "Lebanese")
        self.assertEqual(data["city"], "Beirut")

    def test_tax_basis_profit_and_loss_leaves_out_non_deductible_expenses(self):
        import ledger_reports
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as folder:
            db = Database(Path(folder) / "c.db"); db.initialize("secret12345")
            base = {"expense_date": "05-03-2026", "currency": "USD", "vat": 0, "expense_account": "6261", "payment_account": "5311"}
            db.add_expense(dict(base, description="Taxi - no invoice", without_vat_subtotal=50, tax_deductible=False), 1)
            db.add_expense(dict(base, description="Office rent", without_vat_subtotal=200, tax_deductible=True), 1)
            def expense_total(**extra):
                lines = ledger_reports._load_lines(db, dict({"profit_loss_only": True}, **extra))
                return sum(float(l["debit"] or 0) - float(l["credit"] or 0) for l in lines if str(l["code"]).startswith("6"))
            self.assertAlmostEqual(expense_total(), 250)
            self.assertAlmostEqual(expense_total(tax_basis=True), 200)
            # the trial balance is never changed by the option
            all_lines = ledger_reports._load_lines(db, {"tax_basis": True})
            self.assertAlmostEqual(sum(float(l["debit"] or 0) - float(l["credit"] or 0) for l in all_lines), 0)
            self.assertAlmostEqual(sum(float(l["debit"] or 0) for l in all_lines if str(l["code"]).startswith("6")), 250)
            db.release()


if __name__ == "__main__":
    unittest.main()
