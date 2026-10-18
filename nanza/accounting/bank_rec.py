"""Bank reconciliation: bank statement lines matched with the book lines of a bank account (511 / 512 / 519 / 53)."""
from __future__ import annotations

import json
import re
from datetime import datetime
from decimal import Decimal

from database import display_date, iso_date, utcnow

ZERO = Decimal("0")


def migrate(db):
    db.execute("""CREATE TABLE IF NOT EXISTS bank_statement_lines (id INTEGER PRIMARY KEY, account_code TEXT NOT NULL, currency TEXT NOT NULL, line_date TEXT NOT NULL,
        description TEXT, reference TEXT, amount TEXT NOT NULL, journal_line_id INTEGER, created_at TEXT NOT NULL)""")
    db.execute("""CREATE TABLE IF NOT EXISTS bank_statement_matches (
        id INTEGER PRIMARY KEY, statement_line_id INTEGER NOT NULL REFERENCES bank_statement_lines(id) ON DELETE CASCADE,
        journal_line_id INTEGER NOT NULL REFERENCES journal_lines(id), amount TEXT NOT NULL,
        created_at TEXT NOT NULL, UNIQUE(statement_line_id,journal_line_id))""")
    db.execute("""INSERT OR IGNORE INTO bank_statement_matches(statement_line_id,journal_line_id,amount,created_at)
        SELECT s.id,s.journal_line_id,CAST(ABS(CAST(s.amount AS REAL)) AS TEXT),s.created_at
        FROM bank_statement_lines s WHERE s.journal_line_id IS NOT NULL
        AND EXISTS(SELECT 1 FROM journal_lines l WHERE l.id=s.journal_line_id)""")


def _d(value):
    try: return Decimal(str(value if value not in (None, "") else 0).replace(",", ""))
    except (ArithmeticError, ValueError, TypeError, KeyError): raise ValueError(f"'{value}' is not an amount")


def add_statement_lines(database, account, currency, rows, user_id):
    """rows: [{date, description, reference, amount}] amount + = money in (deposit), - = money out."""
    account = str(account or "").split(" - ", 1)[0].strip(); currency = str(currency or "USD").upper()
    if not account.startswith(("511", "512", "519", "53")): raise ValueError("Choose a bank or cash account (511, 512, 519 or 53)")
    added = 0
    with database.connect() as db:
        for row in rows:
            amount = _d(row.get("amount"))
            if not amount: continue
            day = iso_date(row.get("date"), "Date"); text = str(row.get("description") or "").strip()
            if db.execute("SELECT 1 FROM bank_statement_lines WHERE account_code=? AND currency=? AND line_date=? AND CAST(amount AS REAL)=? AND COALESCE(description,'')=?",
                          (account, currency, day, float(amount), text)).fetchone():
                continue  # 2.9.87: the same statement imported twice is not doubled
            db.execute("INSERT INTO bank_statement_lines(account_code,currency,line_date,description,reference,amount,created_at) VALUES(?,?,?,?,?,?,?)",
                (account, currency, iso_date(row.get("date"), "Date"), str(row.get("description") or "").strip(), str(row.get("reference") or "").strip(), str(amount), utcnow())); added += 1
        db.execute("INSERT INTO audit_log(user_id,action,entity,details,created_at) VALUES(?,?,?,?,?)", (user_id, "import", "bank_statement",
                json.dumps({"account": account, "lines": added}), utcnow()))
    return added


_STATEMENT_DATE = re.compile(r"^\s*(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{2,4})\b")
_STATEMENT_MONEY = re.compile(r"(?<![\d/.\-])-?\d{1,3}(?:,\d{3})*\.\d{2}(?![\d/])(?:\s*(?:CR|DR|-))?", re.I)
_DEBIT_WORDS = re.compile(r"(?i)\b(?:chq|cheque|check|withdraw\w*|charges?|commission|fees?|debit|payment|paid|transfer\s+to|atm|pos|card|tax)\b")


def read_statement_pdf(path):
    """2.9.87: a bank statement in PDF (text, or a scan read with the local OCR). Each line that starts with a date and
    has amounts is a movement; the running balance column gives its sign (balance up = money in, down = money out);
    opening / closing balance lines are not movements. Returns the same rows as the Excel import, plus the balances read."""
    import pdf_import
    from pypdf import PdfReader
    reader = PdfReader(str(path)); texts = []
    for index, page in enumerate(reader.pages):
        text = pdf_import._page_text(page, True) or pdf_import._page_text(page)
        if len(text.strip()) < 20 or text.count("\x1f") > 5:
            try: text = pdf_import._ocr_pdf_pages(path, [index])[0]
            except Exception as exc:
                if not texts: raise ValueError(f"This PDF is a scan and the local OCR could not read it ({exc})")
        texts.append(text)
    rows, opening, closing, previous = [], None, None, None
    for raw in "\n".join(pdf_import.normalize_invoice_text(t) for t in texts).splitlines():
        match = _STATEMENT_DATE.match(raw)
        if not match: continue
        day, month, year = match.groups(); year = int(year) + (2000 if len(year) == 2 else 0)
        try: date = datetime(year, int(month), int(day)).strftime("%Y-%m-%d")
        except ValueError: continue
        rest = raw[match.end():]
        rest = re.sub(r"^\s*\d{1,2}[/.\-]\d{1,2}[/.\-]\d{2,4}", "", rest)  # value date
        amounts = []
        for m in _STATEMENT_MONEY.finditer(rest):
            token = m.group(0).strip(); negative = token.startswith("-") or token.upper().endswith(("DR", "-"))
            value = Decimal(re.sub(r"(?i)[^\d.]", "", token)); amounts.append((-value if negative else value, m.start()))
        description = re.sub(r"\s{2,}", " ", _STATEMENT_MONEY.sub(" ", rest)).strip(" |-")
        low = description.casefold()
        if re.search(r"opening|brought forward|b/f|previous balance|رصيد سابق|solde (?:initial|precedent)", low):
            if amounts: opening = previous = amounts[-1][0]
            continue
        if re.search(r"closing|carried forward|c/f|new balance|رصيد نهائي|solde final", low):
            if amounts: closing = amounts[-1][0]
            continue
        if not amounts: continue
        reference = ""
        ref = re.search(r"(?i)\b(?:chq|cheque|ref|no\.?)\s*#?\s*(\d{3,})", description)
        if ref: reference = ref.group(1)
        if len(amounts) >= 2 and previous is not None:
            balance = amounts[-1][0]; amount = balance - previous; previous = balance
            if abs(abs(amount) - abs(amounts[-2][0])) > Decimal("0.01"):  # the balance column does not follow: trust the amount and its column
                amount = amounts[-2][0] if not _DEBIT_WORDS.search(description) else -abs(amounts[-2][0])
        elif len(amounts) >= 2:
            amount = amounts[0][0] if not _DEBIT_WORDS.search(description) else -abs(amounts[0][0]); previous = amounts[-1][0]
        else:
            amount = -abs(amounts[0][0]) if _DEBIT_WORDS.search(description) else amounts[0][0]
            if previous is not None: previous += amount
        if amount: rows.append({"date": date, "description": description[:200], "reference": reference, "amount": amount})
    if not rows: raise ValueError("No movement was found in this PDF statement (lines must start with a date and show the amounts)")
    return rows if opening is None and closing is None else _with_balances(rows, opening, closing)


class _Rows(list):
    """Rows of a statement, with the opening and closing balances printed on it (when found)."""


def _with_balances(rows, opening, closing):
    result = _Rows(rows); result.opening = opening; result.closing = closing; return result


def read_statement_file(path):
    """Excel / CSV bank statement: Date, Description, Reference and either Amount or Debit / Credit (or Withdrawal / Deposit).
    2.9.87: a PDF statement is read too."""
    if str(path).lower().endswith(".pdf"): return read_statement_pdf(path)
    from importer import _header_map, _date
    if str(path).lower().endswith(".csv"):
        import csv
        with open(path, newline="", encoding="utf-8-sig") as handle: table = [row for row in csv.reader(handle)]
        class Sheet:
            def iter_rows(self, min_row=1, max_row=None, values_only=True):
                return [tuple(r) for r in table[min_row - 1:max_row]]
        sheet = Sheet(); close = lambda: None
    else:
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True); sheet = book.worksheets[0]; close = book.close
    try:
        header, columns = _header_map(sheet, {"date": ("date", "تاريخ", "value date"), "description": ("description", "details", "narration", "بيان", "libell"),
            "reference": ("reference", "ref", "cheque", "check"), "amount": ("amount", "montant", "مبلغ"), "debit": ("debit", "withdrawal", "out", "مدين"),
            "credit": ("credit", "deposit", "in ", "دائن")})
        rows = []
        for row in sheet.iter_rows(min_row=header + 1, values_only=True):
            get = lambda f: row[columns[f]] if f in columns and columns[f] < len(row) else None
            if not any(v not in (None, "") for v in row) or get("date") in (None, ""): continue
            if "amount" in columns: amount = _d(get("amount"))
            else: amount = _d(get("credit")) - _d(get("debit"))  # bank statement: credit = money in, debit = money out
            rows.append({"date": _date(get("date")), "description": get("description") or "", "reference": get("reference") or "", "amount": amount})
        return rows
    finally: close()


def statement_lines(database, account, date_from, date_to):
    with database.connect() as db:
        rows = [dict(r) for r in db.execute("""SELECT s.*,
            COALESCE((SELECT DSUM(CAST(m.amount AS REAL)) FROM bank_statement_matches m WHERE m.statement_line_id=s.id),0) matched_amount
            FROM bank_statement_lines s WHERE s.account_code=? AND s.line_date BETWEEN ? AND ? ORDER BY s.line_date,s.id""",
            (account, date_from, date_to))]
    for row in rows:
        row["remaining_amount"] = max(ZERO, abs(_d(row["amount"])) - _d(row["matched_amount"]))
        row["is_matched"] = row["remaining_amount"] < Decimal("0.005")
    return rows


def book_lines(database, account, currency, date_from, date_to):
    """Book movements of the bank account: debit = money in (+), credit = money out (-)."""
    with database.connect() as db:
        matched = {r["journal_line_id"]: Decimal(str(r["matched_amount"] or 0)) for r in db.execute(
            "SELECT journal_line_id,DSUM(CAST(amount AS REAL)) matched_amount FROM bank_statement_matches GROUP BY journal_line_id")}
        rows = [dict(r) for r in db.execute("""SELECT l.id,e.entry_number,e.entry_date,COALESCE(l.description,e.description) description,e.description entry_description,
            e.source_type,COALESCE(e.voucher_type,'') voucher_type,
            CAST(l.debit AS REAL) debit,CAST(l.credit AS REAL) credit,COALESCE(NULLIF(l.line_currency,''),e.currency) currency,
            CASE WHEN COALESCE(NULLIF(l.line_currency,''),e.currency)<>e.currency AND l.amount IS NOT NULL THEN CAST(l.amount AS REAL) ELSE NULL END line_amount
            FROM journal_lines l JOIN journal_entries e ON e.id=l.entry_id JOIN accounts a ON a.id=l.account_id WHERE a.code=?""", (account,))]
    result = []
    for row in rows:
        try: day = iso_date(row["entry_date"])
        except ValueError: continue
        if not date_from <= day <= date_to or row["currency"] != currency: continue
        sign = 1 if row["debit"] else -1
        value = Decimal(str(row["line_amount"])) * sign if row["line_amount"] is not None else _d(row["debit"]) - _d(row["credit"])
        matched_amount = matched.get(row["id"], ZERO)
        remaining = max(ZERO, abs(value) - matched_amount)
        result.append({**row, "iso_date": day, "amount": value, "matched_amount": matched_amount,
                       "remaining_amount": remaining, "statement_id": row["id"] if remaining < Decimal("0.005") else None})
    return sorted(result, key=lambda r: (r["iso_date"], r["id"]))


def auto_match(database, account, currency, date_from, date_to, days=5, user_id=None):
    """One-to-one: same amount, dates within `days` days, the closest date first."""
    statement = [s for s in statement_lines(database, account, date_from, date_to) if not s["journal_line_id"] and s["currency"] == currency]
    books = [b for b in book_lines(database, account, currency, date_from, date_to) if not b["statement_id"]]
    used = set(); matched = 0
    with database.connect() as db:
        for line in statement:
            day = datetime.strptime(line["line_date"], "%Y-%m-%d")
            amount = _d(line["remaining_amount"]) * (1 if _d(line["amount"]) >= 0 else -1)
            candidates = [b for b in books if b["id"] not in used and abs(b["remaining_amount"] - abs(amount)) < Decimal("0.005")
                          and abs((datetime.strptime(b["iso_date"], "%Y-%m-%d") - day).days) <= days]
            if not candidates: continue
            best = min(candidates, key=lambda b: abs((datetime.strptime(b["iso_date"], "%Y-%m-%d") - day).days))
            used.add(best["id"]); matched += 1
            match_partial(database, line["id"], [{"journal_line_id": best["id"], "amount": str(abs(amount))}], user_id)
    return matched


def match(database, statement_id, journal_line_id, user_id=None):
    """Match the full amount, requiring the statement and book movements to agree."""
    with database.connect() as db:
        line = db.execute("SELECT * FROM bank_statement_lines WHERE id=?", (int(statement_id),)).fetchone()
        if not line: raise KeyError("Statement line not found")
        book = db.execute("""SELECT l.debit,l.credit FROM journal_lines l
            WHERE l.id=?""", (int(journal_line_id),)).fetchone()
        if not book: raise KeyError("Book line not found")
        statement_amount = abs(_d(line["amount"]))
        book_amount = abs(_d(book["debit"]) - _d(book["credit"]))
        if abs(statement_amount - book_amount) >= Decimal("0.005"):
            raise ValueError("Full match requires equal statement and book amounts; use partial allocation instead")
    return match_partial(database, statement_id,
                         [{"journal_line_id": journal_line_id, "amount": str(abs(_d(line["amount"])))}], user_id)


def match_partial(database, statement_id, allocations, user_id=None):
    """Allocate part or all of a statement movement across one or more book lines."""
    statement_id = int(statement_id)
    if not isinstance(allocations, list) or not allocations:
        raise ValueError("Provide at least one book-line allocation")
    with database.connect() as db:
        statement = db.execute("SELECT * FROM bank_statement_lines WHERE id=?", (statement_id,)).fetchone()
        if not statement: raise KeyError("Statement line not found")
        statement_amount = _d(statement["amount"])
        already = db.execute("SELECT COALESCE(DSUM(CAST(amount AS REAL)),0) total FROM bank_statement_matches WHERE statement_line_id=?",
                             (statement_id,)).fetchone()["total"]
        statement_remaining = abs(statement_amount) - Decimal(str(already or 0))
        cleaned = []
        seen = set()
        for allocation in allocations:
            line_id = int(allocation.get("journal_line_id"))
            amount = _d(allocation.get("amount"))
            if line_id in seen: raise ValueError("A book line may appear only once in an allocation")
            seen.add(line_id)
            if db.execute("SELECT 1 FROM bank_statement_matches WHERE statement_line_id=? AND journal_line_id=?",
                          (statement_id, line_id)).fetchone():
                raise ValueError(f"Book line {line_id} is already allocated to this statement line")
            if amount <= 0: raise ValueError("Allocation amounts must be greater than zero")
            row = db.execute("""SELECT l.id,l.debit,l.credit,COALESCE(NULLIF(l.line_currency,''),e.currency) currency,a.code
                FROM journal_lines l JOIN journal_entries e ON e.id=l.entry_id JOIN accounts a ON a.id=l.account_id WHERE l.id=?""",
                (line_id,)).fetchone()
            if not row: raise KeyError(f"Book line {line_id} not found")
            book_amount = _d(row["debit"]) - _d(row["credit"])
            if (book_amount > 0) != (statement_amount > 0):
                raise ValueError("Statement and book movements must have the same direction")
            if row["code"] != statement["account_code"] or row["currency"] != statement["currency"]:
                raise ValueError("Book line must use the same bank account and currency as the statement")
            used = db.execute("SELECT COALESCE(DSUM(CAST(amount AS REAL)),0) total FROM bank_statement_matches WHERE journal_line_id=?",
                              (line_id,)).fetchone()["total"]
            if amount > abs(book_amount) - Decimal(str(used or 0)):
                raise ValueError(f"Allocation exceeds the remaining book amount for line {line_id}")
            cleaned.append((line_id, amount))
        total = sum((amount for _line_id, amount in cleaned), ZERO)
        if total > statement_remaining:
            raise ValueError("Allocations exceed the remaining statement amount")
        for line_id, amount in cleaned:
            db.execute("""INSERT INTO bank_statement_matches(statement_line_id,journal_line_id,amount,created_at)
                VALUES(?,?,?,?)""", (statement_id, line_id, str(amount), utcnow()))
        matched = db.execute("SELECT COALESCE(DSUM(CAST(amount AS REAL)),0) total FROM bank_statement_matches WHERE statement_line_id=?",
                             (statement_id,)).fetchone()["total"]
        if abs(_d(matched) - abs(statement_amount)) < Decimal("0.005") and len(cleaned) == 1:
            db.execute("UPDATE bank_statement_lines SET journal_line_id=? WHERE id=?", (cleaned[0][0], statement_id))
        else:
            db.execute("UPDATE bank_statement_lines SET journal_line_id=NULL WHERE id=?", (statement_id,))
        db.execute("INSERT INTO audit_log(user_id,action,entity,entity_id,details,created_at) VALUES(?,?,?,?,?,?)",
                   (user_id, "match", "bank_statement_line", statement_id,
                    json.dumps({"allocations": [{"journal_line_id": line_id, "amount": str(amount)}
                                                for line_id, amount in cleaned]}), utcnow()))
    return True


def unmatch(database, statement_id):
    with database.connect() as db:
        db.execute("DELETE FROM bank_statement_matches WHERE statement_line_id=?", (int(statement_id),))
        db.execute("UPDATE bank_statement_lines SET journal_line_id=NULL WHERE id=?", (int(statement_id),))
    return True


def delete_statement_line(database, statement_id):
    with database.connect() as db: db.execute("DELETE FROM bank_statement_lines WHERE id=?", (int(statement_id),))
    return True


def post_statement_line(database, statement_id, account_code, user_id):
    """Bank charges, interest, transfers seen only on the statement: book them and match them."""
    with database.connect() as db: line = db.execute("SELECT * FROM bank_statement_lines WHERE id=?", (int(statement_id),)).fetchone()
    if not line: raise KeyError("Statement line not found")
    if line["journal_line_id"]: raise ValueError("This statement line is already matched")
    amount = _d(line["amount"]); other = str(account_code or "").split(" - ", 1)[0].strip()
    if not other: raise ValueError("Choose the account (for example 6278 bank charges)")
    bank_side, other_side = ("D", "C") if amount > 0 else ("C", "D")
    voucher = database.save_journal_voucher({"entry_date": display_date(line["line_date"]), "description": f"Bank: {line['description'] or 'statement line'}",
            "currency": line["currency"], "voucher_type": "03"},
        [{"account_code": line["account_code"], "line_currency": line["currency"], "side": bank_side, "amount": str(abs(amount)), "reference": line["reference"], "description": line["description"]},
         {"account_code": other, "line_currency": line["currency"], "side": other_side, "amount": str(abs(amount)), "reference": line["reference"], "description": line["description"]}], user_id)
    with database.connect() as db:
        bank_line = db.execute("""SELECT l.id FROM journal_lines l JOIN accounts a ON a.id=l.account_id JOIN journal_entries e ON e.id=l.entry_id
            WHERE e.entry_number=? AND a.code=? ORDER BY l.id LIMIT 1""", (voucher["voucher"]["entry_number"], line["account_code"])).fetchone()
    match_partial(database,statement_id,[{"journal_line_id":bank_line["id"],"amount":str(abs(amount))}],user_id)
    return voucher["voucher"]["entry_number"]


def reconciliation(database, account, currency, date_from, date_to, statement_balance=None):
    """Bank balance per statement = book balance + payments not yet cleared - deposits not yet cleared + statement items not booked."""
    books_all = book_lines(database, account, currency, "0000-01-01", date_to)
    statement = statement_lines(database, account, date_from, date_to)
    book_balance = sum((b["amount"] for b in books_all), ZERO)
    # 2.9.87: an opening balance brought forward is not a deposit waiting for the bank (it is the statement's opening
    # balance); book items of earlier months still not matched stay outstanding once earlier statements were imported.
    with database.connect() as db:
        earlier = db.execute("SELECT 1 FROM bank_statement_lines WHERE account_code=? AND currency=? AND line_date<? LIMIT 1", (account, currency, date_from)).fetchone()
    def opening(b): return b.get("source_type") in ("opening", "year_close") or b.get("voucher_type") == "04"
    books = [b for b in books_all if (b["iso_date"] >= date_from or earlier) and not opening(b)]
    outstanding = [{**b, "amount": (Decimal(1) if b["amount"] >= 0 else Decimal(-1)) * b["remaining_amount"]}
                   for b in books if b["remaining_amount"] >= Decimal("0.005")]
    not_booked = [{**s, "amount": (Decimal(1) if _d(s["amount"]) >= 0 else Decimal(-1)) * s["remaining_amount"]}
                  for s in statement if s["remaining_amount"] >= Decimal("0.005") and s["currency"] == currency]
    deposits = sum((b["amount"] for b in outstanding if b["amount"] > 0), ZERO); payments = -sum((b["amount"] for b in outstanding if b["amount"] < 0), ZERO)
    bank_only = sum((_d(s["amount"]) for s in not_booked), ZERO)
    expected = book_balance - deposits + payments + bank_only
    difference = (_d(statement_balance) - expected) if statement_balance not in (None, "") else None
    rows = [["Balance per books at " + display_date(date_to), book_balance.quantize(Decimal("0.01"))],
            ["Less: deposits in the books not yet on the statement", (-deposits).quantize(Decimal("0.01"))],
            ["Add: payments in the books not yet cleared by the bank", payments.quantize(Decimal("0.01"))],
            ["Add / less: statement items not yet booked (charges, interest...)", bank_only.quantize(Decimal("0.01"))],
            ["EXPECTED BALANCE PER BANK STATEMENT", expected.quantize(Decimal("0.01"))]]
    if difference is not None: rows += [["Balance per bank statement (entered)", _d(statement_balance).quantize(Decimal("0.01"))], ["DIFFERENCE (should be 0)", difference.quantize(Decimal("0.01"))]]
    sections = [{"heading": f"Bank reconciliation - account {account} ({currency})", "headers": ["Item", f"Amount ({currency})"], "rows": rows,
            "total_rows": [4] + ([6] if difference is not None else [])},
                {"heading": "Book items not yet on the statement", "headers": ["Date", "Voucher", "Description", "Amount"],
                 "rows": [[display_date(b["iso_date"]), b["entry_number"], b["description"] or "",
                         b["amount"].quantize(Decimal("0.01"))] for b in outstanding] or [["None", "", "", ""]], "total_rows": []},
                {"heading": "Statement items not yet booked", "headers": ["Date", "Reference", "Description", "Amount"],
                 "rows": [[display_date(s["line_date"]), s["reference"] or "", s["description"] or "",
                         _d(s["amount"]).quantize(Decimal("0.01"))] for s in not_booked] or [["None", "", "", ""]], "total_rows": []}]
    company = database.settings()
    return {"title": "Bank Reconciliation", "meta": [f"Company: {company.get('company_name') or '-'}   Account {account}   Period {display_date(date_from)} to {display_date(date_to)}"],
            "sections": sections, "difference": float(difference) if difference is not None else None, "expected": float(expected), "book_balance": float(book_balance),
            "matched": sum(1 for s in statement if s["matched_amount"] > ZERO), "statement_count": len(statement)}
